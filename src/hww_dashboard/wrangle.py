"""Turn the Healthy Working Wales spreadsheets into tidy CSV tables.

Run from the project folder:

    python wrangle.py              (uses config.toml next to wrangle.py)
    python wrangle.py my.toml      (uses another config file)

``python -m hww_dashboard.wrangle [config.toml]`` works the same way,
with config.toml in the current folder as the default.

Settings live in config.toml and the lookup tables in mappings/. Every
run writes the tables to the output folder plus data_quality_report.md.
"""

import hashlib
import json
import logging
import os
import re
import shutil
import sys
import tempfile
import tomllib
import traceback
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from difflib import get_close_matches
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from hww_dashboard.logs import (
    close_logging,
    configure_logging,
    log_format,
    new_run_id,
    prune_logs,
)
from hww_dashboard.output_schema import check_cells, parse_csv
from hww_dashboard.run_record import (
    RECORD_NAME,
    Record,
    file_facts,
    sha256_bytes,
    sha256_file,
    start_record,
    utc_iso,
    write_record,
)

logger = logging.getLogger(__name__)

#: The config used when none is given: config.toml in the current folder.
DEFAULT_CONFIG = Path("config.toml")
UNKNOWN = "Unknown"
OTHER = "Other"
SOURCES = ("signups", "survey", "was", "peer")
COMPLETED_RULES = ("summary_1_sent", "session_1_held", "last_session_held")
#: Where the tables go: CSV files in a folder, or BigQuery tables.
OUTPUT_TARGETS = ("folder", "bigquery")
#: ``${NAME}`` in config text, replaced by an environment variable.
CONFIG_VARIABLE = re.compile(r"\$\{([A-Z_][A-Z0-9_]*)\}")
DATE_IN_NAME = re.compile(r"(\d{4})-(\d{2})-(\d{2})")

MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}

# Column order of every output table. Empty tables are still written
# with these headers.
COLUMNS = {
    "dim_employer": [
        "employer_id",
        "employer_name",
        "sector",
        "industry",
        "size_band",
        "sites_band",
        "in_signups",
        "in_survey",
        "in_was",
        "in_peer",
        "attribute_conflicts",
    ],
    # Job roles and the flag columns' raw text are not output (data
    # minimisation, DECISIONS_en.md 2026-10-18).
    "fact_signup": [
        "signup_id",
        "employer_id",
        "signup_month",
        "survey_flag",
        "survey_dates",
        "was_flag",
        "was_dates",
    ],
    "fact_engagement": [
        "engagement_id",
        "employer_id",
        "date",
        "date_precision",
        "event_type",
        "source",
        "source_ref",
    ],
    "dim_survey_response": [
        "response_id",
        "employer_id",
        "date",
        "sector",
        "industry",
        "size_band",
        "sites_band",
    ],
    "fact_survey_answer": [
        "response_id",
        "employer_id",
        "date",
        "question_id",
        "question",
        "answer_id",
        "answer",
    ],
    "survey_other_text": ["response_id", "question_id", "other_text"],
    "question_catalogue": [
        "question_id",
        "question",
        "question_short",
        "question_type",
        "answer_id",
        "answer",
        "answer_short",
        "answer_order",
    ],
    "fact_was_session": [
        "session_id",
        "was_request_id",
        "employer_id",
        "session_number",
        "session_date",
        "duration",
        "summary_sent_date",
        "topics",
    ],
    "fact_was_session_topic": ["session_id", "topic"],
    "fact_event": [
        "event_id",
        "event_name",
        "event_type",
        "event_date",
        "employer_id",
        "action",
        "action_date",
    ],
    "badges": ["employer_id", "badge", "awarded_date"],
    "dim_date": [
        "date",
        "year",
        "month",
        "month_name",
        "month_start",
        "fin_year",
        "fin_quarter",
    ],
}

# Settings every run needs, by config.toml section. Checked before any
# file is read.
REQUIRED_SETTINGS = {
    "paths": ["input_folder", "output_folder", "mappings_folder"],
    "signups": [
        "sheet",
        "header_row",
        "month",
        "org",
        "role",
        "size",
        "sector",
        "industry",
        "sites",
        "completed_survey",
        "requested_was",
    ],
    "survey": ["sheet", "header_row", "keep_other_free_text"],
    "was": [
        "sheet",
        "header_row",
        "id",
        "org",
        "enquiry_date",
        "fallback_enquiry_date",
        "size",
        "sector",
        "industry",
        "sites",
        "topic_separator",
        "completed_rule",
        "sessions",
    ],
    "peer": [
        "sheet",
        "header_row",
        "id",
        "org",
        "interest_date",
        "size",
        "sector",
        "industry",
    ],
    "employers": ["source_priority"],
}

# The columns each mapping CSV must have.
MAPPING_COLUMNS = {
    "lookup": ["raw_value", "agreed_value"],
    "org_overrides": ["raw_name", "canonical_name"],
    "survey_columns": ["column", "role", "question_id", "option_or_note"],
    "survey_questions": [
        "question_id",
        "question",
        "question_short",
        "question_type",
    ],
    "survey_answers": [
        "question_id",
        "answer",
        "answer_short",
        "answer_order",
    ],
    "survey_answer_translations": ["question_id", "raw_answer", "answer"],
    "agreed_values": ["dimension", "agreed_value"],
}

LOOKUP_NAMES = (
    "sector",
    "industry",
    "size_band",
    "sites_band",
    "was_duration",
    "was_topic",
)
# Survey columns the pipeline reads by role (from
# mappings/survey_columns.csv).
SURVEY_ROLES = (
    "response_id",
    "date",
    "org",
    "sector",
    "industry",
    "size",
    "sites",
)

ASSUMPTIONS = [
    "Sign-up date is the first day of the sign-up month "
    "(date_precision = month); the sheet only records the month.",
    "Survey completions come from the survey export and WAS requests from "
    "the WAS database. The 'Completed survey?' and 'Requested WAS?' "
    "columns on Account signups are kept in fact_signup and only used for "
    "the reconciliation section below.",
    "A month written without a year (e.g. 'Y (Dec)') gets the year closest "
    "to the sign-up month (Dec next to a Jan 2026 sign-up = Dec 2025).",
    "'Y' with no month gives flag = True and no date.",
    "'other_options' is treated as the answer 'Other' (sector and "
    "industry).",
    "Size bands are reduced to 0-9 / 10-49 / 50-249 / 250-999 / 1000+ so "
    "every file fits. WAS '50-250' and '251-999' are treated as 50-249 and "
    "250-999.",
    "Site bands are reduced to 1 / 2-5 / 6+ / Not applicable (remote). WAS "
    "'5-8' is treated as 6+ (overlaps 2-5 at exactly 5 sites).",
    "When rows for one organisation disagree, the most common value wins; "
    "on a tie a specific value beats 'Other', then source priority from "
    "config.toml, then the latest row.",
    "employer_id is a short hash of the lower-case organisation name, so it "
    "stays the same between runs.",
    "A WAS session only counts as 'was_session_held' if it has a date.",
    "Survey % for a question = respondents who ticked the answer / "
    "respondents who gave any answer to that question.",
    "Survey free text ('Other' boxes, 'What parts could be improved?') is "
    "not output unless keep_other_free_text = true in config.toml.",
    "Peer mentoring 'relationships' = expressions of interest (no matching "
    "dates recorded yet).",
    "Job roles, and the text typed into 'Completed survey?' and "
    "'Requested WAS?', are not output; fact_signup keeps only the flags "
    "and dates read from that text.",
]

Config = dict[str, Any]  # one section of config.toml
Lookup = dict[str, str]  # lower-case raw value -> agreed value


class InputError(Exception):
    """A problem in config.toml, a mapping file or an input spreadsheet.

    A person must fix it: the message says which file and what to change.
    ``python wrangle.py`` prints it without a traceback and exits with
    code 1.
    """


# ---------------------------------------------------------------- report


#: Things a person should act on after a run, in the order the report
#: lists them: key -> (what was found, what to do about it).
ATTENTION = {
    "unmapped": (
        "value(s) not in the mapping files",
        "add each to the matching CSV in mappings/ (see Unmapped values)",
    ),
    "new_columns": (
        "new survey column(s)",
        "add each to mappings/survey_columns.csv (see Survey checks)",
    ),
    "missing_columns": (
        "survey column(s) missing from the export",
        "check whether the form changed, and update "
        "mappings/survey_columns.csv (see Survey checks)",
    ),
    "new_answers": (
        "survey answer(s) not in survey_answers.csv",
        "add each option, or its Welsh translation (see Survey checks)",
    ),
    "flag_problems": (
        "unreadable flag(s) on Account signups",
        "correct the cell in the workbook (see Free-text flags)",
    ),
    "shrunk_sources": (
        "source(s) with far fewer rows than the last run",
        "check the upload is complete (see Warnings)",
    ),
}


class Report:
    """Collects notes for data_quality_report.md, grouped by section.

    Attributes:
        sections: Report lines for each section name in ``ORDER``.
        unmapped: Counts of raw values missing from a lookup, keyed by
            (dimension, source, raw value).
        attention: Counts of other things to act on, keyed by the names
            in ``ATTENTION``.
        inputs: One entry per input file read, for the run record: its
            source, file name, size, time, checksum, sheet and rows.
    """

    ORDER = [
        "Inputs",
        "Outputs",
        "Rows dropped",
        "Duplicates",
        "Attribute conflicts",
        "Unmapped values",
        "Survey checks",
        "Free-text flags",
        "Warnings",
        "Reconciliation",
    ]

    def __init__(self) -> None:
        """Start an empty report with every section present."""
        self.sections: dict[str, list[str]] = {name: [] for name in self.ORDER}
        self.unmapped: Counter[tuple[str, str, str]] = Counter()
        self.attention: Counter[str] = Counter()
        self.inputs: list[dict[str, Any]] = []

    def needs(self, item: str, count: int = 1) -> None:
        """Count something a person should act on.

        Args:
            item: A key of ``ATTENTION``.
            count: How many were found (0 does nothing).
        """
        if count:
            self.attention[item] += count

    def to_act_on(self) -> dict[str, int]:
        """List what needs attention, in the order of ``ATTENTION``.

        Returns:
            Key of ``ATTENTION`` -> count, for the non-zero ones. Unmapped
            values count once per distinct raw value.
        """
        counts = Counter(self.attention)
        counts["unmapped"] = len(self.unmapped)
        return {key: counts[key] for key in ATTENTION if counts[key]}

    def add(self, section: str, line: str) -> None:
        """Append one markdown line to a section.

        Args:
            section: One of the names in ``ORDER``.
            line: Markdown text, usually starting with ``- `` for a bullet.
        """
        self.sections[section].append(line)

    def dropped(self, source: str, count: int, reason: str) -> None:
        """Note that rows were removed from a source.

        Does nothing when count is 0.

        Args:
            source: Human-readable source name, e.g. ``"WAS database"``.
            count: Number of rows dropped.
            reason: Why they were dropped.
        """
        if count:
            self.add(
                "Rows dropped",
                f"- {source}: {count} row(s) dropped - {reason}",
            )

    def write(self, path: Path, run_at: datetime) -> None:
        """Write the report as markdown, with unmapped values and assumptions.

        Args:
            path: File to write (overwritten).
            run_at: When the run started, shown at the top of the report.
        """
        text = self.render(run_at, path.name)
        path.write_text(text, encoding="utf-8", newline="\n")

    def render(
        self, run_at: datetime, filename: str = "data_quality_report.md"
    ) -> str:
        """Render the report as markdown, checked for contact details.

        Args:
            run_at: When the run started, shown at the top of the report.
            filename: The report's name, for the contact-detail message.

        Returns:
            The report's text.

        Raises:
            InputError: The report would hold a contact identifier.
        """
        lines = [
            "# Data quality report",
            "",
            f"Run at {run_at:%Y-%m-%d %H:%M}.",
            "",
        ]
        sections = deepcopy(self.sections)
        if self.unmapped:
            table = pd.DataFrame(
                [
                    (d, s, v, n)
                    for (d, s, v), n in sorted(self.unmapped.items())
                ],
                columns=["dimension", "source", "raw value", "rows"],
            )
            sections["Unmapped values"].append(
                "These were set to Unknown (bands) or kept as-is "
                "(topics, durations).",
            )
            sections["Unmapped values"].append(
                "Add them to the matching CSV in mappings/ to fix.",
            )
            sections["Unmapped values"].append("")
            sections["Unmapped values"].append(markdown_table(table))
        for name in self.ORDER:
            lines.append(f"## {name}")
            lines.append("")
            lines.extend(sections[name] or ["None."])
            lines.append("")
        lines.append("## Assumptions")
        lines.append("")
        lines.extend(f"- {a}" for a in ASSUMPTIONS)
        lines.append("")
        text = "\n".join(lines)
        check_public_text(text, filename)
        return text


def markdown_table(df: pd.DataFrame) -> str:
    """Render a small DataFrame as a markdown table.

    Args:
        df: Table to render. Pipe characters in cells become ``/``.

    Returns:
        The markdown text, or ``"None."`` for an empty table.
    """
    if df.empty:
        return "None."
    header = "| " + " | ".join(str(c) for c in df.columns) + " |"
    rule = "|" + "---|" * len(df.columns)
    body = [
        "| " + " | ".join(str(v).replace("|", "/") for v in row) + " |"
        for row in df.itertuples(index=False)
    ]
    return "\n".join([header, rule, *body])


def count_values(values: pd.Series) -> pd.Series:
    """Count each value: most common first, ties in order of appearance.

    pandas' own ``value_counts`` leaves the order of tied counts to the
    pandas version (2.x and 3.x differ), which would make the report
    differ between installs. Grouping keeps values in the order they
    first appear, and a stable sort keeps that order within each count.

    Args:
        values: The values to count. Blanks (NaN) are not counted.

    Returns:
        The number of rows for each value, indexed by value.
    """
    counts = values.groupby(values, sort=False).size()
    return counts.sort_values(ascending=False, kind="stable")


# ---------------------------------------------------------------- cleaning


def clean_text(value: object) -> str:
    """Turn any cell into tidy text.

    Blanks become '', 1 and '1' both become '1', and stray newlines,
    Excel '_x000D_' codes and double spaces are removed.

    Args:
        value: A cell value of any type (text, number, date, None, NaN).

    Returns:
        The tidied text.
    """
    if value is None or value is pd.NA:
        return ""
    if isinstance(value, float):
        if pd.isna(value):
            return ""
        if value.is_integer():
            value = int(value)
    if value is pd.NaT:
        return ""
    text = str(value).replace("_x000D_", " ")
    return " ".join(text.split())


def to_date(value: object) -> date | None:
    """Read a cell as a date.

    Excel dates and serials use the Windows 1900 epoch. ISO text is
    year-first; other text is day-first (04/12/2025 is 4 December).

    Args:
        value: A cell value of any type.

    Returns:
        The date, or None if the cell is blank or not a date.
    """
    if value is pd.NaT or value is pd.NA or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            parsed_excel = from_excel(value)
        except (ValueError, OverflowError):
            return None
        return (
            parsed_excel.date() if isinstance(parsed_excel, datetime) else None
        )
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])", text):
        try:
            return datetime.fromisoformat(text).date()
        except ValueError:
            return None
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    return None if pd.isna(parsed) else parsed.date()


def make_unique(names: Iterable[str]) -> list[str]:
    """Give repeated column names a suffix the same way pandas does.

    Args:
        names: Column names in sheet order.

    Returns:
        The names with repeats numbered: ``Other, Other.1, Other.2 ...``
    """
    names = list(names)
    reserved = set(names)
    used: set[str] = set()
    seen: Counter[str] = Counter()
    result = []
    for name in names:
        candidate = name
        while candidate in used:
            seen[name] += 1
            candidate = f"{name}.{seen[name]}"
            while candidate in reserved:
                seen[name] += 1
                candidate = f"{name}.{seen[name]}"
        used.add(candidate)
        result.append(candidate)
    return result


def suggest(missing: str, choices: Iterable[str]) -> str:
    """Return ' Did you mean ...?' for a misspelt name.

    Args:
        missing: The name that was not found.
        choices: The names that do exist.

    Returns:
        Text to append to an error message, or '' when nothing is close.
    """
    close = get_close_matches(
        missing, [str(c) for c in choices], n=3, cutoff=0.6
    )
    return (
        f" Did you mean {' or '.join(repr(c) for c in close)}?"
        if close
        else ""
    )


# ---------------------------------------------------------------- reading


def load_config(path: Path) -> Config:
    """Read config.toml and check every required setting is present.

    Args:
        path: The config file.

    Returns:
        The whole config as nested dicts, with every ``${NAME}`` replaced
        by the environment variable NAME.

    Raises:
        InputError: The file is missing, is not valid TOML, uses an
            environment variable that isn't set, or lacks a setting.
    """
    try:
        with open(path, "rb") as f:
            cfg = tomllib.load(f)
    except FileNotFoundError:
        raise InputError(f"Config file not found: {path}") from None
    except tomllib.TOMLDecodeError as error:
        raise InputError(f"{path.name} is not valid TOML: {error}") from None
    missing: set[str] = set()
    cfg = dict(interpolate(cfg, missing))
    if missing:
        raise InputError(
            f"{path.name} uses environment variable(s) that aren't set: "
            f"{sorted(missing)}. Set them, or write the value in the config."
        )
    check_config(cfg, path.name)
    return cfg


def interpolate(value: Any, missing: set[str]) -> Any:
    """Replace ``${NAME}`` in config text with environment variables.

    The syntax is py-common's: NAME is upper-case letters, digits and
    underscores, and the replacement is done once, as the config loads.

    Args:
        value: A config value: text, a list or a table (dict).
        missing: Collects the names of variables that aren't set; their
            placeholders are left as they are.

    Returns:
        The value with every placeholder in it replaced.
    """
    if isinstance(value, str):

        def replace(match: re.Match[str]) -> str:
            """Look up one variable.

            Args:
                match: The ``${NAME}`` found.

            Returns:
                The variable's value, or the placeholder if it isn't set.
            """
            name = match.group(1)
            if name not in os.environ:
                missing.add(name)
                return match.group(0)
            return os.environ[name]

        return CONFIG_VARIABLE.sub(replace, value)
    if isinstance(value, list):
        return [interpolate(item, missing) for item in value]
    if isinstance(value, dict):
        return {key: interpolate(item, missing) for key, item in value.items()}
    return value


def check_config(cfg: Config, name: str = "config.toml") -> None:
    """Check config.toml has every setting, with valid values.

    Stops early, with a clear message, before any input file is read.

    Args:
        cfg: The parsed config.
        name: Config file name, for messages.

    Raises:
        InputError: A section or setting is missing or invalid.
    """
    for section, keys in REQUIRED_SETTINGS.items():
        if section not in cfg:
            raise InputError(f"{name}: section [{section}] is missing.")
        if not isinstance(cfg[section], dict):
            raise InputError(f"{name}: [{section}] must be a table.")
        missing = [k for k in keys if k not in cfg[section]]
        if missing:
            raise InputError(
                f"{name}: [{section}] is missing setting(s) {missing}."
            )
    for section in ("signups", "survey", "was", "peer"):
        if "file" not in cfg[section] and "file_pattern" not in cfg[section]:
            raise InputError(
                f"{name}: [{section}] needs 'file' or 'file_pattern'."
            )
    for section, keys in REQUIRED_SETTINGS.items():
        for key in keys:
            value = cfg[section][key]
            valid = True
            if key == "header_row":
                valid = type(value) is int and value >= 1
            elif key == "sheet":
                valid = (type(value) is int and value >= 0) or (
                    isinstance(value, str) and bool(value.strip())
                )
            elif key == "keep_other_free_text":
                valid = type(value) is bool
            elif key in ("sessions", "source_priority"):
                valid = isinstance(value, list)
            else:
                valid = isinstance(value, str) and bool(value.strip())
            if not valid:
                raise InputError(
                    f"{name}: [{section}] {key} has an invalid value. "
                    "Use the type and nonblank value in the shipped config."
                )
    for section in SOURCES:
        for key in ("file", "file_pattern"):
            if key in cfg[section]:
                value = cfg[section][key]
                if not isinstance(value, str) or not value.strip():
                    raise InputError(
                        f"{name}: [{section}] {key} must be nonblank text."
                    )
    folder = cfg["paths"].get("log_folder", "logs")
    if not isinstance(folder, str) or not folder.strip():
        raise InputError(f"{name}: [paths] log_folder must be nonblank text.")
    output = cfg.get("output", {})
    if not isinstance(output, dict):
        raise InputError(f"{name}: [output] must be a table.")
    target = output.get("target", "folder")
    if target not in OUTPUT_TARGETS:
        raise InputError(
            f"{name}: [output] target must be one of "
            f"{list(OUTPUT_TARGETS)}, got {target!r}."
        )
    if target == "bigquery":
        for key in ("project", "dataset"):
            value = output.get(key)
            if not isinstance(value, str) or not value.strip():
                raise InputError(
                    f"{name}: [output] {key} must be nonblank text for "
                    "BigQuery output."
                )
    location = output.get("location", "unset")
    if not isinstance(location, str) or not location.strip():
        raise InputError(f"{name}: [output] location must be nonblank text.")
    for source in SOURCES:
        if type(cfg[source].get("allow_empty", False)) is not bool:
            raise InputError(
                f"{name}: [{source}] allow_empty must be true or false."
            )
    checks = cfg.get("checks", {})
    drop = checks.get("max_row_drop", 0.1) if isinstance(checks, dict) else -1
    if type(drop) not in (int, float) or not 0 <= drop <= 1:
        raise InputError(
            f"{name}: [checks] max_row_drop must be a number from 0 to 1."
        )
    privacy = cfg.get("privacy", {})
    if not isinstance(privacy, dict) or not isinstance(
        privacy.get("personal_columns", []), list
    ):
        raise InputError(
            f"{name}: [privacy] personal_columns must be a list of names."
        )
    personal_columns = privacy.get("personal_columns", [])
    if not all(isinstance(c, str) and c.strip() for c in personal_columns):
        raise InputError(
            f"{name}: [privacy] personal_columns must contain nonblank names."
        )
    personal = {clean_text(c).casefold() for c in personal_columns}
    for source in SOURCES:
        for key, value in cfg[source].items():
            if key not in (
                "file",
                "file_pattern",
                "completed_rule",
                "topic_separator",
            ):
                if (
                    isinstance(value, str)
                    and clean_text(value).casefold() in personal
                ):
                    raise InputError(
                        f"{name}: [{source}] {key} selects a personal column. "
                        "Choose a non-personal business field."
                    )
    numbers: set[int] = set()
    for number, session in enumerate(cfg["was"]["sessions"], start=1):
        if not isinstance(session, dict):
            raise InputError(
                f"{name}: [[was.sessions]] entry {number} must be a table."
            )
        missing = [k for k in ("number", "date") if k not in session]
        if missing:
            raise InputError(
                f"{name}: [[was.sessions]] entry {number} is missing "
                f"{missing}."
            )
        session_number = session["number"]
        if (
            type(session_number) is not int
            or session_number < 1
            or session_number in numbers
        ):
            raise InputError(
                f"{name}: [[was.sessions]] number must be a unique "
                "positive whole number."
            )
        numbers.add(session_number)
        for key in ("date", "duration", "summary_sent", "topics"):
            if key in session and (
                not isinstance(session[key], str) or not session[key].strip()
            ):
                raise InputError(
                    f"{name}: [[was.sessions]] {key} must be nonblank text."
                )
    rule = cfg["was"]["completed_rule"]
    if rule not in COMPLETED_RULES:
        raise InputError(
            f"{name}: [was] completed_rule '{rule}' is not one of "
            f"{list(COMPLETED_RULES)}."
        )
    priority = cfg["employers"]["source_priority"]
    if not all(isinstance(p, str) for p in priority) or (
        sorted(priority) != sorted(SOURCES)
    ):
        raise InputError(
            f"{name}: [employers] source_priority must list each of "
            f"{list(SOURCES)} once, got {priority}."
        )


def date_in_name(path: Path) -> date | None:
    """Find a YYYY-MM-DD date in a file name.

    Args:
        path: The file.

    Returns:
        The last valid date in the name, or None.
    """
    found = None
    for match in DATE_IN_NAME.finditer(path.stem):
        try:
            found = date(*(int(part) for part in match.groups()))
        except ValueError:
            continue
    return found


def newest_match(folder: Path, pattern: str, source: str) -> Path:
    """Pick the newest file in a folder whose name matches a pattern.

    Newest means the latest YYYY-MM-DD date in the file name (exports are
    named like ``...-survey-tool-2026-09-23.xlsx``). Files without a date
    in the name count as older than any dated file. Ties are broken by
    the time the file was last saved. Excel's temporary lock files
    (``~$...``) are ignored.

    Args:
        folder: Folder to search (not searched recursively).
        pattern: Glob pattern, e.g. ``*employer-survey-tool-*.xlsx``.
        source: Name of the source, for messages.

    Returns:
        The chosen file.

    Raises:
        InputError: No file matches.
    """
    matches = [
        p
        for p in folder.glob(pattern)
        if p.is_file() and not p.name.startswith("~$")
    ]
    if not matches:
        present = sorted(p.name for p in folder.glob("*.xlsx"))
        raise InputError(
            f"{source}: no file matching '{pattern}' in {folder}. "
            f"Excel files there: {present or 'none'}. "
            "Check file_pattern in config.toml."
        )

    def newness(path: Path) -> tuple[date, float, str, str]:
        """Rank exports by date, save time and filename.

        Args:
            path: Matching export file.

        Returns:
            A total ordering independent of directory iteration order.
        """
        return (
            date_in_name(path) or date.min,
            path.stat().st_mtime,
            path.name.casefold(),
            path.name,
        )

    chosen = max(matches, key=newness)
    others = sorted(p.name for p in matches if p != chosen)
    logger.info(
        "%s: using %s (newest of %d matching file(s))",
        source,
        chosen.name,
        len(matches),
    )
    if others:
        logger.info("%s: ignoring older file(s) %s", source, others)
    return chosen


def input_file(folder: Path, cfg: Config, source: str) -> Path:
    """Find the spreadsheet for one source.

    Args:
        folder: The input folder.
        cfg: The source's section of config.toml. Uses ``file_pattern``
            if set (newest match wins), otherwise ``file``.
        source: Name of the source, for messages.

    Returns:
        Path to the file.

    Raises:
        InputError: The file does not exist.
    """
    if "file_pattern" in cfg:
        return newest_match(folder, cfg["file_pattern"], source)
    path = folder / str(cfg["file"])
    if not path.is_file():
        raise InputError(
            f"{source}: file not found: {path}."
            f"{suggest(cfg['file'], files_in(folder))} "
            "Check 'file' in config.toml and input_folder in [paths]."
        )
    return path


def files_in(folder: Path) -> list[str]:
    """List the file names in a folder (empty if the folder is missing).

    Args:
        folder: The folder.

    Returns:
        The names of the files and folders directly inside it.
    """
    return [p.name for p in folder.iterdir()] if folder.is_dir() else []


def read_sheet(
    path: Path,
    sheet: str | int,
    header_row: int,
    source: str = "",
    required: Sequence[str] = (),
) -> pd.DataFrame:
    """Read one Excel sheet as text-like objects with tidy, unique headers.

    Blank headers are named after their Excel column letter (e.g. 'B').
    An 'excel_row' column records where each row came from.

    Args:
        path: The .xlsx file.
        sheet: Sheet name, or 0-based position (0 = first sheet).
        header_row: Excel row number (1-based) that holds the column
            headings.
        source: Name of the source, for messages. Defaults to the file
            name.
        required: Column headings that must be present.

    Returns:
        One row per spreadsheet row below the headings. Every cell keeps
        its Excel type (text, number, datetime or NaN); nothing is
        converted yet.

    Raises:
        InputError: The file cannot be opened, the sheet or heading row
            does not exist, or a required column is missing.
    """
    source = source or path.name
    try:
        with pd.ExcelFile(path) as workbook:
            names = [str(n) for n in workbook.sheet_names]
            if isinstance(sheet, int):
                if not 0 <= sheet < len(names):
                    raise InputError(
                        f"{source}: {path.name} has no sheet number {sheet} "
                        f"(0 = first; it has {len(names)}: {names})."
                    )
                sheet_name = names[sheet]
            elif sheet not in names:
                raise InputError(
                    f"{source}: {path.name} has no sheet '{sheet}'."
                    f"{suggest(sheet, names)} Sheets: {names}."
                )
            else:
                sheet_name = sheet
            raw = workbook.parse(sheet_name=sheet, header=None, dtype=object)
    except FileNotFoundError:
        raise InputError(f"{source}: file not found: {path}") from None
    except PermissionError:
        raise InputError(
            f"{source}: cannot open {path} " "(is it open in another program?)"
        ) from None
    except (zipfile.BadZipFile, ValueError) as error:
        raise InputError(
            f"{source}: {path.name} is not a readable .xlsx file "
            f"({error})."
        ) from None

    if not isinstance(header_row, int) or header_row < 1:
        raise InputError(
            f"{source}: header_row must be a whole number from 1, "
            f"got {header_row!r}."
        )
    if header_row > len(raw):
        raise InputError(
            f"{source}: header_row is {header_row} but sheet "
            f"'{sheet_name}' of {path.name} has only {len(raw)} row(s)."
        )
    headers = [
        clean_text(h) or get_column_letter(i + 1)
        for i, h in enumerate(raw.iloc[header_row - 1])
    ]
    df = raw.iloc[header_row:].copy()
    df.columns = pd.Index(make_unique(headers))
    require(
        df,
        required,
        source,
        f"{path.name}, sheet '{sheet_name}', heading row {header_row}",
    )
    df["excel_row"] = range(header_row + 1, header_row + 1 + len(df))
    logger.info(
        "%s: read %d row(s) from %s (sheet '%s')",
        source,
        len(df),
        path.name,
        sheet_name,
    )
    table = df.reset_index(drop=True)
    table.attrs["sheet_name"] = str(sheet_name)
    return table


def note_input(
    report: Report, source: str, path: Path, df: pd.DataFrame, detail: str
) -> None:
    """Record an input file in the report and for the run record.

    Args:
        report: Data quality report.
        source: Name of the source, e.g. ``"WAS database"``.
        path: The file read.
        df: The sheet as returned by ``read_sheet``.
        detail: Words that follow the row count in the report line,
            e.g. ``" (sheet 'Account signups')"``.
    """
    facts = file_facts(path)
    sheet = str(df.attrs.get("sheet_name", ""))
    report.inputs.append({"source": source, "sheet": sheet, "rows": len(df)})
    report.inputs[-1].update(facts)
    report.add(
        "Inputs",
        f"- {source}: {len(df)} rows read{detail} from {path.name} "
        f"(SHA-256 {facts['sha256'][:12]}).",
    )


def require_rows(
    df: pd.DataFrame, cfg: Config, source: str, section: str, path: Path
) -> None:
    """Stop when a source has no usable rows, unless that is allowed.

    An empty upload usually means a broken export, and publishing it
    would empty the dashboard's tables without a word. py-common rejects
    empty sheets for the same reason.

    Args:
        df: The source's rows, after unusable ones were dropped.
        cfg: The source's config section; ``allow_empty = true`` lets
            an empty source through.
        source: Name of the source, for the message.
        section: The config section's name, for the message.
        path: The file read.

    Raises:
        InputError: The source has no rows and allow_empty isn't true.
    """
    if df.empty and cfg.get("allow_empty", False) is not True:
        raise InputError(
            f"{source}: {path.name} has no data rows. An empty upload "
            "usually means the export went wrong, so check the file. If "
            f"no rows is right, set allow_empty = true in [{section}]."
        )


def compare_with_previous(
    report: Report, previous: Record | None, max_drop: float
) -> None:
    """Warn when a source has noticeably fewer rows than the last run.

    The workbooks are cumulative, so their rows should only grow; a
    big drop usually means an incomplete upload. Edits can remove rows
    legitimately, so this warns rather than stops.

    Args:
        report: Data quality report, with this run's inputs.
        previous: The previous published run's record, or None.
        max_drop: The largest share of rows (0 to 1) that may go before
            a warning, e.g. 0.1 for 10%.
    """
    if not previous:
        return
    before = {
        str(entry.get("source")): entry
        for entry in previous.get("inputs", [])
        if isinstance(entry, dict)
    }
    for entry in report.inputs:
        old = before.get(str(entry["source"]))
        old_rows = old.get("rows") if old else None
        if not isinstance(old_rows, int) or old_rows <= 0:
            continue
        drop = 1 - entry["rows"] / old_rows
        if drop > max_drop:
            report.add(
                "Warnings",
                f"- {entry['source']}: {entry['rows']} rows read, "
                f"{drop:.0%} fewer than the previous run ({old_rows}, run "
                f"{previous.get('run_id', 'unknown')}). These workbooks "
                "only grow, so check the upload is complete.",
            )
            report.needs("shrunk_sources")


def read_record(path: Path) -> Record | None:
    """Read the run record published with the previous outputs.

    Args:
        path: ``run_record.json`` in the output folder.

    Returns:
        The record, or None if there is none or it can't be read.
    """
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return None
    except (OSError, ValueError):
        logger.warning("Cannot read the previous run record %s", path)
        return None
    return data if isinstance(data, dict) else None


def require(
    df: pd.DataFrame, columns: Iterable[str], source: str, where: str = ""
) -> None:
    """Stop with a clear message if expected columns are missing.

    Args:
        df: The table that should have the columns.
        columns: Column names that must be present.
        source: Name of the source, for the message.
        where: Extra detail for the message, e.g. file and sheet.

    Raises:
        InputError: At least one column is missing.
    """
    missing = [c for c in columns if c not in df.columns]
    if missing:
        hints = "".join(suggest(c, df.columns) for c in missing)
        place = f" ({where})" if where else ""
        raise InputError(
            f"{source}: column(s) not found{place}: {missing}.{hints} "
            "Check config.toml (or mappings/survey_columns.csv for the "
            "survey)."
        )


def drop_blank_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Remove rows where every cell (apart from excel_row) is blank.

    Args:
        df: A table from ``read_sheet``.

    Returns:
        A copy without the blank rows.
    """
    cells = df.drop(columns="excel_row").map(clean_text)
    return df[(cells != "").any(axis=1)].copy()


def read_csv_text(path: Path, required: Sequence[str] = ()) -> pd.DataFrame:
    """Read a mapping CSV with every value as text.

    Excel may save the file with a byte-order mark; it is dropped.

    Args:
        path: The CSV file.
        required: Column headings the file must have.

    Returns:
        The table. Blank cells are ``''``, never NaN.

    Raises:
        InputError: The file is missing or lacks a required column.
    """
    try:
        table = pd.read_csv(
            path, dtype=str, keep_default_na=False, encoding="utf-8-sig"
        )
    except FileNotFoundError:
        raise InputError(f"Mapping file not found: {path}") from None
    except (pd.errors.EmptyDataError, pd.errors.ParserError, UnicodeError):
        raise InputError(
            f"{path}: save a valid UTF-8 CSV with a header and "
            "the same number of fields in every row."
        ) from None
    except PermissionError:
        raise InputError(
            f"{path}: cannot read mapping file. Close it in other "
            "programs and check read permissions."
        ) from None
    require(table, required, f"mappings/{path.name}")
    return table


def load_lookup(path: Path) -> Lookup:
    """Read a raw_value -> agreed_value CSV.

    Args:
        path: The mapping CSV.

    Returns:
        A dict keyed by tidied, lower-case raw value.
    """
    table = read_csv_text(path, MAPPING_COLUMNS["lookup"])
    return mapping_pairs(table, "raw_value", "agreed_value", path)


def mapping_pairs(
    table: pd.DataFrame, raw: str, agreed: str, path: Path
) -> Lookup:
    """Validate and collect a two-column mapping.

    Args:
        table: CSV records in file order.
        raw: Key column, matched after cleaning and case folding.
        agreed: Target column, retaining display case.
        path: Mapping file for diagnostic messages.

    Returns:
        Unambiguous mapping; identical duplicate rows are harmless.

    Raises:
        InputError: A key or target is blank, or targets conflict.
    """
    result: Lookup = {}
    for number, (key, value) in enumerate(
        zip(table[raw], table[agreed], strict=True), start=2
    ):
        key, value = clean_text(key).casefold(), clean_text(value)
        if not key or not value or (key in result and result[key] != value):
            raise InputError(
                f"{path}: row {number} has a blank or conflicting mapping. "
                "Fill both columns and give each raw value one target."
            )
        result[key] = value
    return result


def check_agreed_values(lookups: Mapping[str, Lookup], path: Path) -> None:
    """Check every mapping target is on the agreed list for its dimension.

    A typo in a mapping file's agreed_value would otherwise create a new
    band without a word.

    Args:
        lookups: From ``load_lookup``, keyed by dimension.
        path: ``mappings/agreed_values.csv``: one row per dimension and
            agreed value, in the order people expect to see them.

    Raises:
        InputError: The list has a blank or unknown dimension, or a
            mapping file targets a value that isn't on the list.
    """
    table = read_csv_text(path, MAPPING_COLUMNS["agreed_values"])
    allowed: dict[str, set[str]] = {}
    for dimension, value in zip(
        table["dimension"], table["agreed_value"], strict=True
    ):
        dimension, value = clean_text(dimension), clean_text(value)
        if dimension not in LOOKUP_NAMES or not value:
            raise InputError(
                f"mappings/{path.name}: every row needs a dimension from "
                f"{list(LOOKUP_NAMES)} and a nonblank agreed_value."
            )
        allowed.setdefault(dimension, set()).add(value)
    for name, lookup in lookups.items():
        unknown = sorted(set(lookup.values()) - allowed.get(name, set()))
        if unknown:
            raise InputError(
                f"mappings/{name}.csv: {unknown} is not an agreed {name} "
                f"value in mappings/{path.name}. Fix the spelling, or add "
                "the value there if it is a new band."
            )


def load_mappings(maps: Path) -> tuple[dict[str, Lookup], Lookup]:
    """Read the band/topic lookups and the organisation-name overrides.

    Args:
        maps: The mappings folder.

    Returns:
        (lookups, overrides): lookups maps a name such as ``"sector"`` to
        its lookup; overrides maps a lower-case raw organisation name to
        the name to use instead.
    """
    lookups = {
        name: load_lookup(maps / f"{name}.csv") for name in LOOKUP_NAMES
    }
    check_agreed_values(lookups, maps / "agreed_values.csv")
    table = read_csv_text(
        maps / "org_overrides.csv", MAPPING_COLUMNS["org_overrides"]
    )
    path = maps / "org_overrides.csv"
    pairs = mapping_pairs(table, "raw_name", "canonical_name", path)
    overrides: Lookup = {}
    for raw, canonical in pairs.items():
        key = name_key(raw)
        if overrides.get(key, canonical) != canonical:
            raise InputError(
                f"{path}: two raw names that are the same organisation "
                "(ignoring case, spacing and look-alike characters) point "
                "at different canonical names. Give them one target."
            )
        overrides[key] = canonical
    resolved: Lookup = {}
    for key in overrides:
        visited: set[str] = set()
        current = key
        while current in overrides:
            if current in visited:
                raise InputError(
                    f"{path}: organisation overrides contain a cycle. "
                    "Point each chain at one canonical name."
                )
            visited.add(current)
            target = overrides[current]
            if name_key(target) == current:
                break
            current = name_key(target)
        resolved[key] = target
    return lookups, resolved


# ---------------------------------------------------------------- mapping


def map_value(
    raw: object,
    lookup: Mapping[str, str],
    dimension: str,
    source: str,
    report: Report,
    keep_unmapped: bool = False,
) -> str:
    """Map a raw value to the agreed list.

    Blank -> Unknown. A value missing from the lookup is logged in the
    report and becomes Unknown, or is kept as-is when keep_unmapped is
    True.

    Args:
        raw: The cell value.
        lookup: From ``load_lookup``.
        dimension: Lookup name, for the report (e.g. ``"sector"``).
        source: Source name, for the report (e.g. ``"was"``).
        report: Where unmapped values are counted.
        keep_unmapped: Keep unknown text instead of returning Unknown
            (used for free-form fields such as WAS topics, where a new
            value is still worth showing).

    Returns:
        The agreed value.
    """
    text = clean_text(raw)
    if not text:
        return "" if keep_unmapped else UNKNOWN
    agreed = lookup.get(text.casefold())
    if agreed:
        return agreed
    report.unmapped[(dimension, source, text)] += 1
    return text if keep_unmapped else UNKNOWN


def map_bands(
    df: pd.DataFrame,
    columns: Mapping[str, str],
    lookups: Mapping[str, Lookup],
    source: str,
    report: Report,
) -> None:
    """Add mapped sector / industry / size / sites columns to a table.

    Args:
        df: The source table (changed in place).
        columns: Output column -> raw column, e.g.
            ``{"size_band": "Number of employees"}``. Output columns are
            added in this order.
        lookups: From ``load_mappings``.
        source: Source name, for the report.
        report: Where unmapped values are counted.
    """
    for dimension, raw_column in columns.items():
        df[dimension] = [
            map_value(v, lookups[dimension], dimension, source, report)
            for v in df[raw_column]
        ]


def org_name(raw: object, overrides: Mapping[str, str]) -> str:
    """Clean an organisation name and apply the manual override file.

    Args:
        raw: The cell value.
        overrides: From ``load_mappings``.

    Returns:
        The name to use ('' if blank).
    """
    name = clean_text(raw)
    return overrides.get(name_key(name), name)


def name_key(name: str) -> str:
    """Give the form in which organisation names are compared and hashed.

    Two names are the same organisation when their keys match. The key
    is the cleaned name with invisible formatting characters removed
    (Unicode category Cf: soft hyphens, zero-width spaces), look-alike
    characters folded by NFKC normalisation (full-width letters,
    ligatures, non-breaking spaces), case-folded, normalised again, and
    with runs of whitespace collapsed to one space. docs/identifiers.md
    gives the rule in SQL, and tests/data/identifier_cases.json the
    expected key and id of awkward names.

    Args:
        name: An organisation name, already cleaned by ``clean_text``.

    Returns:
        The key, e.g. ``"org 11"``.
    """
    visible = "".join(c for c in name if unicodedata.category(c) != "Cf")
    folded = unicodedata.normalize("NFKC", visible).casefold()
    return " ".join(unicodedata.normalize("NFKC", folded).split())


def employer_id_for(name: str) -> str:
    """Make a stable id from an organisation name: E and 8 hex digits.

    Args:
        name: Organisation name, already cleaned.

    Returns:
        'E' followed by the first 8 hex digits of the MD5 of the name's
        ``name_key``.
    """
    digest = hashlib.md5(name_key(name).encode("utf-8")).hexdigest()
    return "E" + digest[:8]


def iso(d: date | None) -> str:
    """Format a date as YYYY-MM-DD, or '' for None.

    Args:
        d: The date.

    Returns:
        The text.
    """
    return d.isoformat() if d else ""


def fin_year(d: date) -> str:
    """Label the UK financial year (April to March) a date falls in.

    Args:
        d: The date.

    Returns:
        E.g. '2025-26' for any date from 1 April 2025 to 31 March 2026.
    """
    start = d.year if d.month >= 4 else d.year - 1
    return f"{start}-{str(start + 1)[2:]}"


# ---------------------------------------------------------------- flags


def nearest_year(month: int, ref_month: date) -> int:
    """Choose the year that puts a month closest to a reference month.

    Args:
        month: Month number, 1-12.
        ref_month: The reference (the sign-up month).

    Returns:
        The year, e.g. 'Dec' near Jan 2026 -> 2025. An exact 6-month tie
        goes to the later year.
    """
    best_year, best_gap = ref_month.year, None
    for year in (ref_month.year - 1, ref_month.year, ref_month.year + 1):
        gap = abs((year - ref_month.year) * 12 + month - ref_month.month)
        # '<=' so that an exact 6-month tie goes to the later year.
        if best_gap is None or gap <= best_gap:
            best_year, best_gap = year, gap
    return best_year


def parse_flag(
    raw: object, ref_month: date | None
) -> tuple[bool | None, list[date], str]:
    """Read text like 'Y', 'Y (Dec)' or 'Y (Nov 25 & Aug 26)'.

    A month without a year gets the year that puts it closest to
    ref_month.

    Args:
        raw: The cell value.
        ref_month: The sign-up month, used to guess missing years. May be
            None.

    Returns:
        (flag, dates, problem):
          flag    True / False, or None if the text could not be read
          dates   list of first-of-month dates found in the brackets
          problem '' or a note for the report
    """
    text = clean_text(raw)
    if not text:
        return False, [], ""
    first_word = text.split("(")[0].strip().casefold()
    if first_word in ("n", "no"):
        return False, [], ""
    if first_word not in ("y", "yes"):
        return None, [], f"could not read '{text}'"

    dates = []
    problem = ""
    inside = re.search(r"\((.*)\)", text)
    if inside:
        for part in re.split(r"&|,| and ", inside.group(1)):
            part = part.strip()
            if not part:
                continue
            match = re.match(r"([A-Za-z]+)\.?\s*(\d{2}|\d{4})?$", part)
            month = (
                MONTHS.get(match.group(1)[:3].casefold()) if match else None
            )
            if match is None or month is None:
                problem = f"could not read date '{part}' in '{text}'"
                continue
            if match.group(2):
                year = int(match.group(2))
                if year < 100:
                    year += 2000
            elif ref_month:
                year = nearest_year(month, ref_month)
            else:
                problem = f"no year and no sign-up month for '{text}'"
                continue
            dates.append(date(year, month, 1))
    return True, dates, problem


# ---------------------------------------------------------------- sources


def load_signups(
    cfg: Config,
    folder: Path,
    lookups: Mapping[str, Lookup],
    overrides: Lookup,
    report: Report,
    today: date,
) -> pd.DataFrame:
    """Read the Account signups tab: one row per user sign-up.

    Args:
        cfg: The [signups] section of config.toml.
        folder: The input folder.
        lookups: From ``load_mappings``.
        overrides: Organisation-name overrides.
        report: Data quality report.
        today: The run date. Dates read from the flag text after it are
            noted in the report.

    Returns:
        The sheet's rows with an organisation name, plus the columns org,
        month, role, sector, industry, size_band, sites_band, signup_id,
        signup_month and survey_/was_ flag, dates and raw.
    """
    keys = [
        "month",
        "org",
        "role",
        "size",
        "sector",
        "industry",
        "sites",
        "completed_survey",
        "requested_was",
    ]
    path = input_file(folder, cfg, "Account signups")
    df = read_sheet(
        path,
        cfg["sheet"],
        cfg["header_row"],
        "Account signups",
        [cfg[k] for k in keys],
    )
    note_input(
        report, "Account signups", path, df, f" (sheet '{cfg['sheet']}')"
    )

    # The month is a merged cell, so only the first row of each month has
    # it.
    raw_months = df[str(cfg["month"])]
    months = raw_months.map(to_date)
    invalid_months = raw_months.map(clean_text).ne("") & months.isna()
    if invalid_months.any():
        rows = df.loc[invalid_months, "excel_row"].tolist()
        raise InputError(
            f"{path.name}: invalid sign-up month in Excel row(s) {rows}. "
            "Use a valid date, or leave merged continuation cells blank."
        )
    df["month"] = months.ffill()

    df["org"] = df[cfg["org"]].map(lambda v: org_name(v, overrides))
    blank = df["org"] == ""
    report.dropped(
        "Account signups",
        int(blank.sum()),
        "no organisation name (blank rows)",
    )
    df = df[~blank].copy()
    require_rows(df, cfg, "Account signups", "signups", path)

    no_month = df["month"].isna()
    if no_month.any():
        report.add(
            "Warnings",
            f"- Account signups: {int(no_month.sum())} row(s) have no month.",
        )

    # Sometimes the organisation name was typed into the role column.
    df["role"] = df[cfg["role"]].map(clean_text)
    role_is_org = df["role"].str.casefold() == df["org"].str.casefold()
    for row in df[role_is_org].itertuples():
        report.add(
            "Warnings",
            f"- Account signups row {row.excel_row}: role was the org name "
            f"'{row.role}', cleared.",
        )
    df.loc[role_is_org, "role"] = ""

    bands = {
        "sector": cfg["sector"],
        "industry": cfg["industry"],
        "size_band": cfg["size"],
        "sites_band": cfg["sites"],
    }
    map_bands(df, bands, lookups, "signups", report)

    df["signup_id"] = df["excel_row"].map(lambda r: f"SU{r:04d}")
    df["signup_month"] = df["month"]

    # Turn the 'Y (Dec)' style text into a flag and dates, keeping the raw
    # text.
    flag_columns = (
        (cfg["completed_survey"], "survey"),
        (cfg["requested_was"], "was"),
    )
    for column, prefix in flag_columns:
        flags, dates, raws = [], [], []
        for value, month, excel_row in zip(
            df[column], df["month"], df["excel_row"], strict=True
        ):
            flag, found, problem = parse_flag(value, month)
            where = f"- Account signups row {excel_row}, '{column}'"
            if problem:
                report.add("Free-text flags", f"{where}: {problem}")
                report.needs("flag_problems")
            for d in found:
                if d > today:
                    report.add(
                        "Free-text flags",
                        f"{where}: '{clean_text(value)}' read as {d}, "
                        "which is in the future.",
                    )
            flags.append(flag)
            dates.append(";".join(d.isoformat() for d in found))
            raws.append(clean_text(value))
        df[f"{prefix}_flag"] = flags
        df[f"{prefix}_dates"] = dates
        df[f"{prefix}_raw"] = raws
        with_text = sum(1 for r in raws if r)
        no_date = sum(
            1 for f, d in zip(flags, dates, strict=True) if f and not d
        )
        report.add(
            "Free-text flags",
            f"- '{column}': {with_text} row(s) filled in, {no_date} of them "
            "say 'Y' with no month.",
        )
    return df


def validate_survey_columns(columns_map: pd.DataFrame) -> None:
    """Reject ambiguous or incomplete survey column specifications.

    Args:
        columns_map: The records from survey_columns.csv.

    Raises:
        InputError: A column is repeated, or a role/question is invalid.
    """
    allowed = {*SURVEY_ROLES, "single", "tick", "other_text", "drop"}
    invalid = (
        ~columns_map["role"].isin(allowed)
        | columns_map["column"].map(clean_text).eq("")
        | columns_map["column"].duplicated(keep=False)
    )
    questions = columns_map["role"].isin(["single", "tick", "other_text"])
    invalid |= questions & columns_map["question_id"].map(clean_text).eq("")
    invalid |= columns_map["role"].eq("tick") & columns_map[
        "option_or_note"
    ].map(clean_text).eq("")
    if invalid.any():
        raise InputError(
            "mappings/survey_columns.csv: invalid column mapping. "
            "Use a known role, unique nonblank column names, question IDs "
            "for answer columns and an option for each tick."
        )


def unpivot_survey(
    df: pd.DataFrame,
    columns_map: pd.DataFrame,
    answer_lookup: Mapping[tuple[str, str], str],
    translations: Mapping[tuple[str, str], str],
    keep_other_text: bool,
) -> tuple[pd.DataFrame, pd.DataFrame, Counter[tuple[str, str]]]:
    """Turn the wide survey export into long rows.

    Args:
        df: The survey export, with a ``response_id`` column added.
        columns_map: survey_columns.csv as a DataFrame (column, role,
            question_id, option_or_note).
        answer_lookup: (question_id, lower-case answer) -> answer as
            written in survey_answers.csv.
        translations: (question_id, lower-case raw answer) -> English
            answer.
        keep_other_text: Keep the free text typed into 'Other' boxes.

    Returns:
        (answers, other_text, new_answers):
          answers      one row per ticked option or single answer:
                       response_id, question_id, answer
          other_text   free text typed into 'Other' boxes (empty unless
                       keep_other_text)
          new_answers  (question_id, answer) pairs not found in
                       survey_answers.csv
    """
    validate_survey_columns(columns_map)
    answer_rows: list[dict[str, str]] = []
    other_rows: list[dict[str, str]] = []
    new_answers: Counter[tuple[str, str]] = Counter()
    question_columns = columns_map[
        columns_map["role"].isin(["single", "tick", "other_text"])
    ]
    question_columns = question_columns[
        question_columns["column"].isin(df.columns)
    ]
    for _, row in df.iterrows():
        response_id = row["response_id"]
        for spec in question_columns.itertuples():
            text = clean_text(row[spec.column])
            if not text:
                continue
            qid = str(spec.question_id)
            if spec.role == "tick":
                # Any value in a tick-box column means the box was ticked;
                # the option comes from the mapping.
                option = clean_text(spec.option_or_note)
                answer = answer_lookup.get((qid, option.casefold()), option)
            elif spec.role == "single":
                answer = translations.get((qid, text.casefold()), text)
                answer = answer_lookup.get((qid, answer.casefold()), answer)
            else:
                if keep_other_text:
                    other_rows.append(
                        {
                            "response_id": response_id,
                            "question_id": qid,
                            "other_text": text,
                        }
                    )
                continue
            if (qid, answer.casefold()) not in answer_lookup:
                new_answers[(qid, answer)] += 1
            answer_rows.append(
                {
                    "response_id": response_id,
                    "question_id": qid,
                    "answer": answer,
                }
            )
    answers = pd.DataFrame(
        answer_rows, columns=["response_id", "question_id", "answer"]
    )
    other_text = pd.DataFrame(other_rows, columns=COLUMNS["survey_other_text"])
    return answers.drop_duplicates(), other_text, new_answers


def load_survey(
    cfg: Config,
    folder: Path,
    maps: Path,
    lookups: Mapping[str, Lookup],
    overrides: Lookup,
    report: Report,
    personal_columns: Sequence[str] = (),
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Read the employer survey export.

    The export has one row per response and one column per question or
    tick box.

    Args:
        cfg: The [survey] section of config.toml.
        folder: The input folder.
        maps: The mappings folder.
        lookups: From ``load_mappings``.
        overrides: Organisation-name overrides.
        report: Data quality report.
        personal_columns: Source headings that may only have the drop role.

    Returns:
        (survey, answers, other_text, columns_map):
          survey       one row per response with response_id, date, org
                       and the bands
          answers      long answers from ``unpivot_survey``
          other_text   'Other' free text (empty unless
                       keep_other_free_text)
          columns_map  survey_columns.csv

    Raises:
        InputError: survey_columns.csv has no column for a needed role,
            or that column is not in the export.
    """
    columns_map = read_csv_text(
        maps / "survey_columns.csv", MAPPING_COLUMNS["survey_columns"]
    )

    personal = {clean_text(c).casefold() for c in personal_columns}
    selected_personal = columns_map["column"].map(
        lambda c: clean_text(c).casefold() in personal
    ) & columns_map["role"].ne("drop")
    if selected_personal.any():
        raise InputError(
            "mappings/survey_columns.csv: personal columns must use "
            "the drop role. Correct the mapping before running again."
        )
    validate_survey_columns(columns_map)

    def column_for(role: str) -> str:
        """Return the export column that survey_columns.csv gives a role.

        Args:
            role: A role in ``SURVEY_ROLES``, e.g. ``"date"``.

        Returns:
            The export column's name.

        Raises:
            InputError: No column has that role.
        """
        found = columns_map.loc[columns_map["role"] == role, "column"]
        if found.empty:
            raise InputError(
                f"Survey: no column with role '{role}' in "
                "mappings/survey_columns.csv."
            )
        if len(found) != 1:
            raise InputError(
                "mappings/survey_columns.csv: give each source role "
                "exactly one column."
            )
        return str(found.iloc[0])

    role_columns = {role: column_for(role) for role in SURVEY_ROLES}
    path = input_file(folder, cfg, "Survey export")
    df = read_sheet(
        path,
        cfg["sheet"],
        cfg["header_row"],
        "Survey export",
        list(role_columns.values()),
    )
    note_input(
        report, "Survey export", path, df, f", {len(df.columns) - 1} columns,"
    )

    known = set(columns_map["column"])
    in_file = [c for c in df.columns if c != "excel_row"]
    unexpected = [c for c in in_file if c not in known]
    missing = [c for c in columns_map["column"] if c not in in_file]
    report.add(
        "Survey checks",
        f"- {len(in_file)} columns in the export, {len(known)} in "
        "mappings/survey_columns.csv.",
    )
    if unexpected:
        report.add(
            "Survey checks",
            "- **New columns not in survey_columns.csv (ignored):** "
            f"{unexpected}",
        )
        logger.warning(
            "Survey export has %d column(s) not in survey_columns.csv",
            len(unexpected),
        )
        report.needs("new_columns", len(unexpected))
    if missing:
        report.add(
            "Survey checks",
            "- **Columns in survey_columns.csv but not in the export:** "
            f"{missing}",
        )
        logger.warning(
            "Survey export lacks %d column(s) in survey_columns.csv",
            len(missing),
        )
        report.needs("missing_columns", len(missing))
    dropped = list(columns_map.loc[columns_map["role"] == "drop", "column"])
    report.add(
        "Survey checks",
        "- Columns never read (personal, technical or free text): "
        f"{', '.join(dropped)}.",
    )

    df = drop_blank_rows(df)
    df["response_id"] = df[role_columns["response_id"]].map(clean_text)
    no_id = df["response_id"] == ""
    report.dropped("Survey export", int(no_id.sum()), "no Entry Id")
    df = df[~no_id]
    repeated = df["response_id"].duplicated()
    report.dropped("Survey export", int(repeated.sum()), "repeated Entry Id")
    df = df[~repeated].copy()
    require_rows(df, cfg, "Survey export", "survey", path)

    df["date"] = df[role_columns["date"]].map(to_date)
    df["org"] = df[role_columns["org"]].map(lambda v: org_name(v, overrides))
    bands = {
        "sector": role_columns["sector"],
        "industry": role_columns["industry"],
        "size_band": role_columns["size"],
        "sites_band": role_columns["sites"],
    }
    map_bands(df, bands, lookups, "survey", report)
    no_org = int((df["org"] == "").sum())
    if no_org:
        report.add(
            "Warnings",
            f"- Survey export: {no_org} response(s) have a blank Company, "
            "so they cannot be linked to an employer (kept with a blank "
            "employer_id).",
        )

    catalogue = read_csv_text(
        maps / "survey_answers.csv", MAPPING_COLUMNS["survey_answers"]
    )
    answer_lookup = {
        (q, a.casefold()): a
        for q, a in zip(
            catalogue["question_id"], catalogue["answer"], strict=True
        )
    }
    trans = read_csv_text(
        maps / "survey_answer_translations.csv",
        MAPPING_COLUMNS["survey_answer_translations"],
    )
    translations: dict[tuple[str, str], str] = {}
    question_ids = set(
        read_csv_text(
            maps / "survey_questions.csv", MAPPING_COLUMNS["survey_questions"]
        )["question_id"]
    )
    if not set(trans["question_id"]).issubset(question_ids):
        raise InputError(
            "mappings/survey_answer_translations.csv: unknown question ID. "
            "Use IDs from survey_questions.csv."
        )
    for qid, group in trans.groupby("question_id", sort=False):
        pairs = mapping_pairs(
            group,
            "raw_answer",
            "answer",
            maps / "survey_answer_translations.csv",
        )
        translations.update(
            {(str(qid), raw): answer for raw, answer in pairs.items()}
        )

    answers, other_text, new_answers = unpivot_survey(
        df,
        columns_map,
        answer_lookup,
        translations,
        cfg["keep_other_free_text"],
    )
    for (qid, answer), n in sorted(new_answers.items()):
        report.add(
            "Survey checks",
            "- Answer not in survey_answers.csv (kept, add it or a "
            f"translation): {qid} '{answer}' x{n}",
        )
    report.needs("new_answers", len(new_answers))
    return df, answers, other_text, columns_map


def load_was(
    cfg: Config,
    folder: Path,
    lookups: Mapping[str, Lookup],
    overrides: Lookup,
    report: Report,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Read the WAS database: one row per support request.

    Each request has up to three sessions.

    Args:
        cfg: The [was] section of config.toml.
        folder: The input folder.
        lookups: From ``load_mappings``.
        overrides: Organisation-name overrides.
        report: Data quality report.

    Returns:
        (was, sessions, topics):
          was       one row per request with was_request_id, org,
                    enquiry_date and the bands
          sessions  one row per session that has any detail (columns as
                    fact_was_session, with org instead of employer_id)
          topics    one row per (session_id, topic)
    """
    keys = [
        "id",
        "org",
        "enquiry_date",
        "fallback_enquiry_date",
        "size",
        "sector",
        "industry",
        "sites",
    ]
    session_columns = [
        s[k]
        for s in cfg["sessions"]
        for k in ("date", "duration", "summary_sent", "topics")
        if k in s
    ]
    path = input_file(folder, cfg, "WAS database")
    df = read_sheet(
        path,
        cfg["sheet"],
        cfg["header_row"],
        "WAS database",
        [cfg[k] for k in keys] + session_columns,
    )
    note_input(report, "WAS database", path, df, "")

    df = drop_blank_rows(df)
    df["org"] = df[cfg["org"]].map(lambda v: org_name(v, overrides))
    blank = df["org"] == ""
    report.dropped("WAS database", int(blank.sum()), "no organisation name")
    df = df[~blank].copy()
    require_rows(df, cfg, "WAS database", "was", path)

    ids = []
    for value, excel_row in zip(df[cfg["id"]], df["excel_row"], strict=True):
        text = clean_text(value)
        if not text:
            text = f"row{excel_row}"
            report.add(
                "Warnings",
                f"- WAS row {excel_row}: no {cfg['id']}, using id '{text}'.",
            )
        ids.append(f"WAS-{text}")
    df["was_request_id"] = ids
    require_unique_ids(df, "was_request_id", path)

    dates = []
    for value, fallback, excel_row in zip(
        df[cfg["enquiry_date"]],
        df[cfg["fallback_enquiry_date"]],
        df["excel_row"],
        strict=True,
    ):
        enquiry = to_date(value)
        if enquiry is None:
            enquiry = to_date(fallback)
            report.add(
                "Warnings",
                f"- WAS row {excel_row}: no {cfg['enquiry_date']}, used "
                f"{cfg['fallback_enquiry_date']} ({iso(enquiry) or 'blank'}).",
            )
        dates.append(enquiry)
    df["enquiry_date"] = dates

    bands = {
        "sector": cfg["sector"],
        "industry": cfg["industry"],
        "size_band": cfg["size"],
        "sites_band": cfg["sites"],
    }
    map_bands(df, bands, lookups, "was", report)

    sessions, topics = unpivot_was_sessions(df, cfg, lookups, report)
    check_was_sessions(df, sessions, report)
    return df, sessions, topics


def unpivot_was_sessions(
    df: pd.DataFrame,
    cfg: Config,
    lookups: Mapping[str, Lookup],
    report: Report,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Turn the session 1-3 column groups into one row per session.

    Also makes one row per session topic. A session with no date,
    duration, summary date or topics is skipped.

    Args:
        df: WAS requests with was_request_id and org.
        cfg: The [was] section of config.toml (sessions,
            topic_separator).
        lookups: From ``load_mappings`` (was_duration, was_topic).
        report: Data quality report.

    Returns:
        (sessions, topics) as described in ``load_was``.
    """
    session_rows: list[dict[str, Any]] = []
    topic_rows: list[dict[str, str]] = []
    for _, row in df.iterrows():
        for s in cfg["sessions"]:
            session_date = to_date(row[s["date"]])
            duration = (
                map_value(
                    row[s["duration"]],
                    lookups["was_duration"],
                    "was_duration",
                    "was",
                    report,
                    True,
                )
                if "duration" in s
                else ""
            )
            summary = (
                to_date(row[s["summary_sent"]])
                if "summary_sent" in s
                else None
            )
            raw_topics = clean_text(row[s["topics"]]) if "topics" in s else ""
            if not (session_date or duration or summary or raw_topics):
                continue
            session_id = f"{row['was_request_id']}-S{s['number']}"
            topic_list: list[str] = []
            for part in raw_topics.split(cfg["topic_separator"]):
                if clean_text(part):
                    topic = map_value(
                        part,
                        lookups["was_topic"],
                        "was_topic",
                        "was",
                        report,
                        True,
                    )
                    if topic not in topic_list:
                        topic_list.append(topic)
            for topic in topic_list:
                topic_rows.append({"session_id": session_id, "topic": topic})
            session_rows.append(
                {
                    "session_id": session_id,
                    "was_request_id": row["was_request_id"],
                    "org": row["org"],
                    "session_number": s["number"],
                    "session_date": session_date,
                    "duration": duration,
                    "summary_sent_date": summary,
                    "topics": "; ".join(topic_list),
                }
            )
    session_columns = [
        "session_id",
        "was_request_id",
        "org",
        "session_number",
        "session_date",
        "duration",
        "summary_sent_date",
        "topics",
    ]
    sessions = pd.DataFrame(session_rows, columns=session_columns)
    topics = pd.DataFrame(
        topic_rows, columns=COLUMNS["fact_was_session_topic"]
    )
    return sessions, topics


def check_was_sessions(
    was: pd.DataFrame, sessions: pd.DataFrame, report: Report
) -> None:
    """Warn about undated sessions and sessions dated before their enquiry.

    Args:
        was: WAS requests with enquiry_date.
        sessions: From ``unpivot_was_sessions``.
        report: Data quality report.
    """
    undated = sessions[sessions["session_date"].isna()]
    if len(undated):
        counts = undated["session_number"].value_counts().sort_index()
        detail = ", ".join(f"session {n}: {c}" for n, c in counts.items())
        report.add(
            "Warnings",
            f"- WAS: {len(undated)} session row(s) have details but no date "
            f"({detail}). Kept in fact_was_session, not counted as "
            "'was_session_held'.",
        )
    if not sessions.empty:
        merged = sessions.merge(
            was[["was_request_id", "enquiry_date"]], on="was_request_id"
        )
        early = merged[
            merged["session_date"].notna()
            & merged["enquiry_date"].notna()
            & (merged["session_date"] < merged["enquiry_date"])
        ]
        if len(early):
            examples = ", ".join(
                f"{r.session_id} ({r.session_date} < {r.enquiry_date})"
                for r in early.head(5).itertuples()
            )
            report.add(
                "Warnings",
                f"- WAS: {len(early)} session(s) are dated before their "
                f"enquiry date, e.g. {examples}. Kept as-is.",
            )


def require_unique_ids(df: pd.DataFrame, column: str, path: Path) -> None:
    """Reject source keys that would duplicate downstream identifiers.

    Args:
        df: Cleaned source records, including generated fallback IDs.
        column: Identifier column to validate.
        path: Source workbook for the error message.

    Raises:
        InputError: Two records have the same identifier.
    """
    if df[column].duplicated().any():
        raise InputError(
            f"{path.name}: duplicate {column}. "
            "Give each retained source row a unique participant ID."
        )


def load_peer(
    cfg: Config,
    folder: Path,
    lookups: Mapping[str, Lookup],
    overrides: Lookup,
    report: Report,
) -> pd.DataFrame:
    """Read the peer mentoring list: one row per expression of interest.

    Args:
        cfg: The [peer] section of config.toml.
        folder: The input folder.
        lookups: From ``load_mappings``.
        overrides: Organisation-name overrides.
        report: Data quality report.

    Returns:
        Rows with an organisation, plus peer_id, org, interest_date and
        the bands (sites_band is always Unknown: this file does not
        collect it).
    """
    keys = ("id", "org", "interest_date", "size", "sector", "industry")
    path = input_file(folder, cfg, "Peer mentoring")
    df = read_sheet(
        path,
        cfg["sheet"],
        cfg["header_row"],
        "Peer mentoring",
        [cfg[k] for k in keys],
    )
    note_input(report, "Peer mentoring", path, df, "")

    df = drop_blank_rows(df)
    df["org"] = df[cfg["org"]].map(lambda v: org_name(v, overrides))
    blank = df["org"] == ""
    report.dropped("Peer mentoring", int(blank.sum()), "no organisation name")
    df = df[~blank].copy()
    require_rows(df, cfg, "Peer mentoring", "peer", path)

    df["peer_id"] = [
        f"PM-{clean_text(v) or 'row' + str(r)}"
        for v, r in zip(df[cfg["id"]], df["excel_row"], strict=True)
    ]
    require_unique_ids(df, "peer_id", path)
    df["interest_date"] = df[cfg["interest_date"]].map(to_date)
    bands = {
        "sector": cfg["sector"],
        "industry": cfg["industry"],
        "size_band": cfg["size"],
    }
    map_bands(df, bands, lookups, "peer", report)
    df["sites_band"] = UNKNOWN  # not collected in this file
    return df


# ---------------------------------------------------------------- employers


def pick_value(
    values: Iterable[str],
    sources: Iterable[str],
    positions: Iterable[int],
    priority: Sequence[str],
) -> tuple[str, list[str]]:
    """Choose one value for an employer from all its rows.

    Most common known value wins. Ties: a specific value beats 'Other',
    then the higher-priority source, then the latest row.

    Args:
        values: The attribute's value on each row (Unknown is ignored).
        sources: The source of each row (``signups``, ``survey``,
            ``was`` or ``peer``).
        positions: Each row's position in the combined table (higher =
            later).
        priority: Source names, most trusted first.

    Returns:
        (chosen value, sorted list of all distinct known values).
    """
    candidates: dict[str, dict[str, int]] = {}
    for value, source, position in zip(
        values, sources, positions, strict=True
    ):
        if value == UNKNOWN:
            continue
        c = candidates.setdefault(
            value, {"count": 0, "source_rank": len(priority), "latest": -1}
        )
        c["count"] += 1
        c["source_rank"] = min(c["source_rank"], priority.index(source))
        c["latest"] = max(c["latest"], position)
    if not candidates:
        return UNKNOWN, []
    ranked = sorted(
        candidates.items(),
        key=lambda item: (
            -item[1]["count"],
            item[0] == OTHER,
            item[1]["source_rank"],
            -item[1]["latest"],
        ),
    )
    return ranked[0][0], sorted(candidates)


def combine_org_rows(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    """Stack the organisation and band columns of every source.

    Args:
        sources: Source name -> cleaned source table, in a fixed order.

    Returns:
        Columns org, sector, industry, size_band, sites_band, source,
        position. ``position`` numbers the rows 0, 1, 2 ... in stacking
        order.
    """
    attrs = ["org", "sector", "industry", "size_band", "sites_band"]
    org_rows = pd.concat(
        [
            frame[attrs].assign(source=source)
            for source, frame in sources.items()
        ],
        ignore_index=True,
    )
    org_rows["position"] = range(len(org_rows))
    return org_rows


def build_dim_employer(
    org_rows: pd.DataFrame, priority: Sequence[str], report: Report
) -> pd.DataFrame:
    """Make one row per organisation, with a single value per attribute.

    Rows from every source are used. Rows whose names differ only by
    case or spacing count as the same organisation.

    Args:
        org_rows: From ``combine_org_rows``.
        priority: Source names, most trusted first (config.toml
            source_priority).
        report: Data quality report.

    Returns:
        dim_employer.

    Raises:
        InputError: Two different names hash to the same employer_id.
    """
    org_rows = org_rows[org_rows["org"] != ""].copy()
    org_rows["key"] = org_rows["org"].map(name_key)
    rows = []
    for _key, group in org_rows.groupby("key", sort=False):
        name = Counter(group["org"]).most_common(1)[0][0]
        spellings = sorted(set(group["org"]))
        if len(spellings) > 1:
            report.add(
                "Duplicates", f"- Spellings merged into '{name}': {spellings}"
            )
        row: dict[str, Any] = {
            "employer_id": employer_id_for(name),
            "employer_name": name,
        }
        conflicts = []
        for attribute in ("sector", "industry", "size_band", "sites_band"):
            value, distinct = pick_value(
                group[attribute], group["source"], group["position"], priority
            )
            row[attribute] = value
            if len(distinct) > 1:
                conflicts.append(f"{attribute}: {' | '.join(distinct)}")
        row["attribute_conflicts"] = "; ".join(conflicts)
        for source in SOURCES:
            row[f"in_{source}"] = bool((group["source"] == source).any())
        rows.append(row)
    dim = pd.DataFrame(rows, columns=COLUMNS["dim_employer"])

    if dim["employer_id"].duplicated().any():
        raise InputError(
            "Two different organisation names produce the same "
            "employer_id. That is very unlikely; ask a developer to "
            "lengthen the id (docs/identifiers.md)."
        )

    with_conflicts = dim[dim["attribute_conflicts"] != ""]
    report.add(
        "Attribute conflicts",
        f"{len(with_conflicts)} employer(s) have rows that disagree. "
        "The chosen value is in dim_employer.",
    )
    report.add("Attribute conflicts", "")
    report.add(
        "Attribute conflicts",
        markdown_table(
            with_conflicts[["employer_name", "attribute_conflicts"]]
        ),
    )
    return dim


def report_duplicates(
    signups: pd.DataFrame, dim_employer: pd.DataFrame, report: Report
) -> None:
    """Summarise repeated sign-up organisations and overlap between sources.

    Args:
        signups: From ``load_signups``.
        dim_employer: From ``build_dim_employer``.
        report: Data quality report.
    """
    repeated = count_values(signups["org"].map(name_key))
    repeated = repeated[repeated > 1]
    report.add(
        "Duplicates",
        f"- Account signups: {len(signups)} sign-up rows for "
        f"{signups['org'].map(name_key).nunique()} organisations.",
    )
    report.add(
        "Duplicates",
        f"- {len(repeated)} organisation(s) appear more than once "
        "(each row kept as a user in fact_signup):",
    )
    for key, n in repeated.items():
        is_key = dim_employer["employer_name"].map(name_key) == key
        name = dim_employer.loc[is_key, "employer_name"].iloc[0]
        report.add("Duplicates", f"  - {name}: {n} rows")
    in_count = dim_employer[
        ["in_signups", "in_survey", "in_was", "in_peer"]
    ].sum(axis=1)
    report.add(
        "Duplicates",
        f"- dim_employer: {len(dim_employer)} employers, "
        f"{int((in_count > 1).sum())} found in more than one source.",
    )
    not_in_signups = dim_employer.loc[
        ~dim_employer["in_signups"], "employer_name"
    ].tolist()
    if not_in_signups:
        report.add(
            "Duplicates",
            "- Employers not on Account signups (added from other files): "
            f"{not_in_signups}",
        )


def attach_employer_ids(
    frames: Iterable[pd.DataFrame], dim_employer: pd.DataFrame
) -> None:
    """Add an employer_id column to each table, matching on the org name.

    The match ignores case. Rows with no organisation get ''. Empty
    tables receive an empty employer_id column.

    Args:
        frames: Tables with an ``org`` column (changed in place).
        dim_employer: From ``build_dim_employer``.
    """
    ids = dict(
        zip(
            dim_employer["employer_name"].map(name_key),
            dim_employer["employer_id"],
            strict=True,
        )
    )
    for frame in frames:
        frame["employer_id"] = frame["org"].map(name_key).map(ids).fillna("")


# ---------------------------------------------------------------- outputs


def build_survey_tables(
    survey: pd.DataFrame,
    answers: pd.DataFrame,
    columns_map: pd.DataFrame,
    dim_employer: pd.DataFrame,
    maps: Path,
    report: Report,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Make dim_survey_response, fact_survey_answer and question_catalogue.

    A response whose own sector/industry/size/sites is Unknown takes its
    employer's value from dim_employer.

    Args:
        survey: From ``load_survey``, with employer_id.
        answers: Long answers from ``load_survey``.
        columns_map: survey_columns.csv.
        dim_employer: From ``build_dim_employer``.
        maps: The mappings folder.
        report: Data quality report.

    Returns:
        (responses, fact_answers, catalogue).
    """
    responses = survey[
        [
            "response_id",
            "employer_id",
            "date",
            "sector",
            "industry",
            "size_band",
            "sites_band",
        ]
    ].copy()
    employer_attrs = dim_employer.set_index("employer_id")
    filled = 0
    for attribute in ("sector", "industry", "size_band", "sites_band"):
        for i in responses.index:
            employer_id = responses.at[i, "employer_id"]
            if responses.at[i, attribute] == UNKNOWN and employer_id:
                responses.at[i, attribute] = employer_attrs.at[
                    employer_id, attribute
                ]
                filled += 1
    if filled:
        report.add(
            "Survey checks",
            f"- {filled} blank response attribute(s) filled from "
            "dim_employer.",
        )

    catalogue = build_question_catalogue(maps, columns_map, answers)
    fact_answers = answers.merge(
        responses[["response_id", "employer_id", "date"]], on="response_id"
    )
    fact_answers = fact_answers.merge(
        catalogue[["question_id", "question", "answer_id", "answer"]],
        on=["question_id", "answer"],
        how="left",
    )
    report.add(
        "Survey checks",
        f"- {len(responses)} response(s) gave {len(fact_answers)} answer "
        "row(s).",
    )
    is_multi = catalogue["question_type"] == "multi"
    multi = catalogue.loc[is_multi, "question_id"].unique()
    for qid in multi:
        n = fact_answers.loc[
            fact_answers["question_id"] == qid, "response_id"
        ].nunique()
        report.add("Survey checks", f"  - {qid}: {n} respondent(s)")
    return responses, fact_answers, catalogue


def build_engagement(
    signups: pd.DataFrame,
    survey: pd.DataFrame,
    was: pd.DataFrame,
    sessions: pd.DataFrame,
    peer: pd.DataFrame,
    completed_rule: str,
    report: Report,
    today: date,
) -> pd.DataFrame:
    """Make one row per engagement event, from every source.

    Args:
        signups: Account signups with employer_id.
        survey: Survey responses with employer_id.
        was: WAS requests with employer_id.
        sessions: WAS sessions with employer_id.
        peer: Peer mentoring interest with employer_id.
        completed_rule: When a WAS request counts as completed (see
            config.toml).
        report: Data quality report.
        today: The run date. Events dated after it are noted in the
            report.

    Returns:
        fact_engagement.

    Raises:
        InputError: completed_rule is not a known rule.
    """
    if completed_rule not in COMPLETED_RULES:
        raise InputError(
            f"Unknown was.completed_rule '{completed_rule}' in " "config.toml."
        )
    events = []

    def add(
        employer_id: object,
        when: object,
        precision: str,
        event_type: str,
        source: str,
        ref: object,
    ) -> None:
        """Record one event. date_precision is blank when there is no date.

        Args:
            employer_id: The employer, or '' when not linked.
            when: The event date, or None.
            precision: ``"day"`` or ``"month"``.
            event_type: One of the event types in ``fact_engagement``.
            source: The source the event came from.
            ref: The source row's own id.
        """
        events.append(
            {
                "engagement_id": f"{event_type}-{ref}",
                "employer_id": employer_id,
                "date": when,
                "date_precision": precision if when else "",
                "event_type": event_type,
                "source": source,
                "source_ref": ref,
            }
        )

    for row in signups.itertuples():
        add(
            row.employer_id,
            row.signup_month,
            "month",
            "account_signup",
            "signups",
            row.signup_id,
        )
    for row in survey.itertuples():
        add(
            row.employer_id,
            row.date,
            "day",
            "survey_completed",
            "survey",
            row.response_id,
        )
    for row in was.itertuples():
        add(
            row.employer_id,
            row.enquiry_date,
            "day",
            "was_request",
            "was",
            row.was_request_id,
        )
    for row in sessions.itertuples():
        if row.session_date:
            add(
                row.employer_id,
                row.session_date,
                "day",
                "was_session_held",
                "was",
                row.session_id,
            )

    # WAS completed, using the rule chosen in config.toml.
    for request_id, group in sessions.groupby("was_request_id"):
        first = group[group["session_number"] == 1]
        if completed_rule == "summary_1_sent":
            when = first["summary_sent_date"].iloc[0] if len(first) else None
        elif completed_rule == "session_1_held":
            when = first["session_date"].iloc[0] if len(first) else None
        elif completed_rule == "last_session_held":
            held = group["session_date"].dropna()
            when = max(held) if len(held) else None
        else:
            raise InputError(
                f"Unknown was.completed_rule '{completed_rule}' in "
                "config.toml."
            )
        if when and not pd.isna(when):
            add(
                group["employer_id"].iloc[0],
                when,
                "day",
                "was_completed",
                "was",
                str(request_id),
            )

    for row in peer.itertuples():
        add(
            row.employer_id,
            row.interest_date,
            "day",
            "peer_mentoring_interest",
            "peer",
            row.peer_id,
        )

    fact = pd.DataFrame(events, columns=COLUMNS["fact_engagement"])
    undated = fact[fact["date"].isna()]
    for event_type, n in count_values(undated["event_type"]).items():
        report.add(
            "Warnings",
            f"- fact_engagement: {n} '{event_type}' event(s) have no date "
            "(counted, but not on monthly charts).",
        )
    future = fact[fact["date"].notna() & (fact["date"] > today)]
    for event_type, n in count_values(future["event_type"]).items():
        report.add(
            "Warnings",
            f"- fact_engagement: {n} '{event_type}' event(s) are dated after "
            "today.",
        )
    return fact


def build_question_catalogue(
    maps: Path, columns_map: pd.DataFrame, answers: pd.DataFrame
) -> pd.DataFrame:
    """Make one row per question and answer option, in form order.

    Tick-box options and answers that turned up in the data but are
    missing from survey_answers.csv are added at the end of their
    question (answer_order 999).

    Args:
        maps: The mappings folder.
        columns_map: survey_columns.csv.
        answers: Long answers from ``unpivot_survey``.

    Returns:
        question_catalogue, with answer_id = question_id + 2-digit
        position.

    Raises:
        InputError: A mapping file is missing, or answer_order is not a
            whole number.
    """
    questions = read_csv_text(
        maps / "survey_questions.csv", MAPPING_COLUMNS["survey_questions"]
    )
    options = read_csv_text(
        maps / "survey_answers.csv", MAPPING_COLUMNS["survey_answers"]
    )
    if (
        questions["question_id"].eq("").any()
        or questions["question_id"].duplicated().any()
        or questions["question"].eq("").any()
        or not questions["question_type"].isin(["single", "multi"]).all()
    ):
        raise InputError(
            "mappings/survey_questions.csv: use unique nonblank question "
            "IDs, nonblank questions and single/multi question types."
        )
    answer_keys = options.assign(answer=options["answer"].str.casefold())
    if (
        options["answer"].eq("").any()
        or answer_keys.duplicated(["question_id", "answer"]).any()
    ):
        raise InputError(
            "mappings/survey_answers.csv: give each question unique, "
            "nonblank answers."
        )
    mapped_ids = columns_map.loc[
        columns_map["role"].isin(["single", "tick", "other_text"]),
        "question_id",
    ]
    known_ids = set(questions["question_id"])
    for filename, ids in (
        ("survey_answers.csv", options["question_id"]),
        ("survey_columns.csv", mapped_ids),
        ("survey_columns.csv", answers["question_id"]),
    ):
        if not set(ids).issubset(known_ids):
            raise InputError(
                f"mappings/{filename}: unknown question ID. "
                "Use IDs defined in survey_questions.csv."
            )
    try:
        options["answer_order"] = options["answer_order"].astype(int)
    except ValueError:
        raise InputError(
            "mappings/survey_answers.csv: every answer_order must be a "
            "whole number."
        ) from None

    known = set(
        zip(
            options["question_id"],
            options["answer"].str.casefold(),
            strict=True,
        )
    )
    extra = []
    ticks = columns_map[columns_map["role"] == "tick"]
    seen = list(
        zip(ticks["question_id"], ticks["option_or_note"], strict=True)
    ) + list(zip(answers["question_id"], answers["answer"], strict=True))
    for qid, answer in seen:
        if (qid, answer.casefold()) not in known:
            known.add((qid, answer.casefold()))
            extra.append(
                {
                    "question_id": qid,
                    "answer": answer,
                    "answer_short": answer,
                    "answer_order": 999,
                }
            )
    if extra:
        options = pd.concat([options, pd.DataFrame(extra)], ignore_index=True)

    catalogue = questions.merge(options, on="question_id", how="left")
    catalogue = catalogue.sort_values(["question_id", "answer_order"])
    catalogue = catalogue.reset_index(drop=True)
    position = catalogue.groupby("question_id").cumcount() + 1
    catalogue["answer_id"] = (
        catalogue["question_id"] + "-" + position.map("{:02d}".format)
    )
    return catalogue[COLUMNS["question_catalogue"]]


def build_dim_date(*date_lists: Iterable[Any]) -> pd.DataFrame:
    """Make a calendar of whole financial years covering every date.

    It runs from the 1 April before the earliest date to the 31 March
    after the latest.

    Args:
        *date_lists: Any number of lists of dates. Blanks (None, NaT)
            are ignored.

    Returns:
        dim_date, one row per day from 1 April to 31 March. Empty if
        there are no dates.
    """
    all_dates = [
        d for dates in date_lists for d in dates if d and not pd.isna(d)
    ]
    if not all_dates:
        return pd.DataFrame(columns=COLUMNS["dim_date"])
    first, last = min(all_dates), max(all_dates)
    start = date(first.year if first.month >= 4 else first.year - 1, 4, 1)
    end = date(last.year + 1 if last.month >= 4 else last.year, 3, 31)
    days = pd.date_range(start, end, freq="D")
    return pd.DataFrame(
        {
            "date": days.date,
            "year": days.year,
            "month": days.month,
            "month_name": days.strftime("%B"),
            "month_start": days.to_period("M").start_time.date,
            "fin_year": [fin_year(d) for d in days.date],
            "fin_quarter": [f"Q{(d.month - 4) % 12 // 3 + 1}" for d in days],
        }
    )


def calendar_dates(
    engagement: pd.DataFrame,
    fact_answers: pd.DataFrame,
    sessions: pd.DataFrame,
    signups: pd.DataFrame,
) -> list[list[Any]]:
    """Collect every date the calendar must cover.

    Args:
        engagement: fact_engagement.
        fact_answers: fact_survey_answer.
        sessions: WAS sessions (session and summary dates).
        signups: Account signups (dates read from the free-text flags).

    Returns:
        Lists of dates for ``build_dim_date``.
    """
    session_dates = (
        list(sessions["session_date"]) + list(sessions["summary_sent_date"])
        if not sessions.empty
        else []
    )
    flag_cells = pd.concat([signups["survey_dates"], signups["was_dates"]])
    flag_dates = [
        date.fromisoformat(d)
        for cell in flag_cells
        for d in cell.split(";")
        if d
    ]
    return [
        list(engagement["date"]),
        list(fact_answers["date"]),
        session_dates,
        flag_dates,
    ]


def survey_percentages(
    answers: pd.DataFrame, question_id: str
) -> pd.DataFrame:
    """Work out Figure 13 style numbers for one question.

    ticks       = respondents who chose the answer
    respondents = respondents who gave any answer to the question
    percent     = ticks / respondents * 100

    Args:
        answers: Long answers (response_id, question_id, answer).
        question_id: E.g. ``"Q04"``.

    Returns:
        One row per answer chosen at least once: answer, ticks,
        respondents, percent; most ticked first.
    """
    rows = answers[answers["question_id"] == question_id]
    respondents = rows["response_id"].nunique()
    result = (
        rows.groupby("answer")["response_id"]
        .nunique()
        .rename("ticks")
        .reset_index()
    )
    result["respondents"] = respondents
    result["percent"] = (
        result["ticks"] / respondents * 100 if respondents else 0.0
    )
    return result.sort_values("ticks", ascending=False).reset_index(drop=True)


def reconcile(
    signups: pd.DataFrame,
    survey: pd.DataFrame,
    was: pd.DataFrame,
    report: Report,
) -> None:
    """Compare the manual flags on Account signups with the real records.

    The real records are the survey responses and the WAS requests.

    Args:
        signups: Account signups with employer_id and the flag columns.
        survey: Survey responses with employer_id.
        was: WAS requests with employer_id.
        report: Data quality report.
    """
    for label, flag_column, other, other_name in (
        ("survey", "survey_flag", survey, "a survey response"),
        ("WAS", "was_flag", was, "a WAS record"),
    ):
        # The column holds True, False and None objects; "== True" gives a
        # clean mask that treats None as not flagged.
        is_flagged = signups[flag_column] == True  # noqa: E712
        flagged = set(signups.loc[is_flagged, "employer_id"])
        actual = set(other["employer_id"]) - {""}
        report.add(
            "Reconciliation",
            f"- {label}: {len(flagged)} employer(s) flagged 'Y' on Account "
            f"signups, {len(actual)} with {other_name}; "
            f"{len(flagged & actual)} in both, "
            f"{len(flagged - actual)} flagged with no record, "
            f"{len(actual - flagged)} with a record but no flag.",
        )


def check_public_text(text: str, filename: str) -> None:
    """Stop recognisable contact identifiers from reaching an output.

    This guard catches email addresses, UK phone numbers (0 or +44 and
    10 more digits, spaces, dashes and brackets allowed), IPv4
    addresses and web links. It cannot identify arbitrary personal
    names in free-form text.

    Args:
        text: Text destined for an output table or report.
        filename: Output name to identify the repair without echoing data.

    Raises:
        InputError: Text contains a recognisable personal identifier.
    """
    pattern = (
        r"[\w.+-]+@[\w-]+\.[\w.]+|"
        r"(?:\+44|\b0)(?:[\s()-]*\d){9,10}\b|"
        r"\b(?:\d{1,3}\.){3}\d{1,3}\b|https?://"
    )
    if re.search(pattern, text, flags=re.IGNORECASE):
        raise InputError(
            f"{filename}: personal identifiers found in retained text. "
            "Remove contact details from source fields and mappings "
            "before running again."
        )


def csv_text(df: pd.DataFrame, name: str) -> str:
    """Render a table as the CSV text that every writer starts from.

    Args:
        df: The table. Missing columns come out empty; extra columns are
            left out.
        name: Table name, a key of ``COLUMNS``.

    Returns:
        The agreed columns as CSV, dates as YYYY-MM-DD, with LF line
        endings on every system, so a run gives the same bytes on
        Windows and Linux.
    """
    table = df.reindex(columns=COLUMNS[name]).copy()
    for column in table.columns:
        table[column] = table[column].map(
            lambda v: v.isoformat() if isinstance(v, (date, datetime)) else v
        )
    return table.to_csv(index=False, lineterminator="\n")


def write_csv(df: pd.DataFrame, folder: Path, name: str) -> int:
    """Write a table with its agreed columns. Dates come out as YYYY-MM-DD.

    The text is checked for contact details and against the table's
    schema (``hww_dashboard/schema/``) before anything is written.

    Args:
        df: The table. Missing columns are written empty; extra columns
            are left out.
        folder: Output folder.
        name: Table name, a key of ``COLUMNS``.

    Returns:
        Number of rows written.

    Raises:
        InputError: The table holds a contact identifier.
        SchemaError: The table breaks its schema (a bug).
    """
    text = csv_text(df, name)
    check_public_text(text, f"{name}.csv")
    check_cells(parse_csv(text), name)
    (folder / f"{name}.csv").write_bytes(text.encode("utf-8"))
    return len(df)


def add_output_counts(
    report: Report, counts: list[tuple[str, int]], engagement: pd.DataFrame
) -> None:
    """Add the rows per table and per event type to the report.

    Args:
        report: Data quality report (the copy about to be published).
        counts: (table name, rows) for each table.
        engagement: fact_engagement, for the event-type summary.
    """
    report.add(
        "Outputs",
        markdown_table(pd.DataFrame(counts, columns=["table", "rows"])),
    )
    report.add("Outputs", "")
    report.add("Outputs", "fact_engagement by event type:")
    report.add("Outputs", "")
    by_type = count_values(engagement["event_type"]).rename_axis("event_type")
    report.add("Outputs", markdown_table(by_type.reset_index(name="rows")))


@dataclass(frozen=True)
class BigQueryTarget:
    """Where, and through what, a run publishes to BigQuery.

    Attributes:
        module: ``hww_dashboard.bigquery_output``, imported only when
            BigQuery output is chosen, so the pipeline runs without the
            bigquery extra.
        client: The BigQuery client.
        dataset: ``project.dataset``.
    """

    module: Any
    client: Any
    dataset: str


def connect_bigquery(settings: Config) -> BigQueryTarget:
    """Get ready to publish to BigQuery, before any file is read.

    Args:
        settings: The [output] section of the config.

    Returns:
        The target to publish to.

    Raises:
        InputError: The bigquery extra isn't installed, or there are no
            Google credentials.
    """
    try:
        from hww_dashboard import bigquery_output
    except ImportError:
        raise InputError(
            "BigQuery output needs the bigquery extra: "
            'pip install -e ".[bigquery]" -c constraints.txt'
        ) from None
    project = str(settings["project"])
    try:
        client = bigquery_output.make_client(project, settings.get("location"))
    except bigquery_output.PublishSetupError as error:
        raise InputError(str(error)) from None
    return BigQueryTarget(
        bigquery_output, client, f"{project}.{settings['dataset']}"
    )


def previous_record(target: BigQueryTarget | None, out: Path) -> Record | None:
    """Fetch the record of the last published run, to compare with.

    Args:
        target: The BigQuery target, or None for a folder.
        out: The output folder (used when target is None).

    Returns:
        The last record, or None if there was no earlier run.

    Raises:
        InputError: BigQuery refused access, or couldn't be reached.
    """
    if target is None:
        return read_record(out / RECORD_NAME)
    try:
        record: Record | None = target.module.latest_record(
            target.dataset, target.client
        )
    except target.module.PublishSetupError as error:
        raise InputError(str(error)) from None
    except OSError as error:
        raise InputError(
            f"Could not reach BigQuery ({type(error).__name__}). Check the "
            "network and run again; nothing was changed."
        ) from None
    return record


def publish_bigquery(
    tables: dict[str, pd.DataFrame],
    engagement: pd.DataFrame,
    report: Report,
    run_at: datetime,
    record: Record,
    target: BigQueryTarget,
) -> list[tuple[str, int]]:
    """Publish every table, the report and the run record to BigQuery.

    The tables pass the same checks as in a folder (contact details,
    schema) and are sent as the same CSV text, so both targets publish
    the same rows. ``tables`` and ``report`` are updated as
    ``write_outputs`` updates them.

    Args:
        tables: Table name -> DataFrame (changed in place).
        engagement: fact_engagement, for the event-type summary.
        report: Data quality report.
        run_at: When the run started, for the report.
        record: The run record so far.
        target: From :func:`connect_bigquery`.

    Returns:
        (table name, rows published) for each table.

    Raises:
        InputError: A table holds a contact identifier, or BigQuery
            reports a problem a person can fix (access, a missing
            dataset, a changed table, no network).
    """
    texts: dict[str, str] = {}
    counts = []
    for name, df in tables.items():
        text = csv_text(df, name)
        check_public_text(text, f"{name}.csv")
        check_cells(parse_csv(text), name)
        texts[name] = text
        counts.append((name, len(df)))
    staged_report = deepcopy(report)
    add_output_counts(staged_report, counts, engagement)
    report_text = staged_report.render(run_at)
    outputs: dict[str, dict[str, Any]] = {
        name: {"rows": n, "sha256": sha256_bytes(texts[name].encode())}
        for name, n in counts
    }
    outputs["data_quality_report"] = {
        "sha256": sha256_bytes(report_text.encode())
    }
    finished = {
        **record,
        "outcome": "published",
        "finished_at": utc_iso(),
        "outputs": outputs,
    }
    try:
        target.module.publish(
            texts, report_text, finished, target.dataset, target.client
        )
    except target.module.PublishSetupError as error:
        raise InputError(str(error)) from None
    except OSError as error:
        raise InputError(
            f"Could not reach BigQuery ({type(error).__name__}). Check the "
            "network and run again; nothing was changed."
        ) from None
    for name, df in tables.items():
        tables[name] = df.reindex(columns=COLUMNS[name])
    report.sections = staged_report.sections
    return counts


def write_outputs(
    tables: dict[str, pd.DataFrame],
    out: Path,
    engagement: pd.DataFrame,
    report: Report,
    run_at: datetime,
    record: Record | None = None,
) -> list[tuple[str, int]]:
    """Write every table and the data quality report.

    Each entry in ``tables`` is replaced by the table as written (agreed
    columns only).

    Args:
        tables: Table name -> DataFrame (changed in place).
        out: Output folder.
        engagement: fact_engagement, for the event-type summary.
        report: Data quality report.
        run_at: When the run started, for the report.
        record: The run record so far. If given, it is completed with
            each output's rows and checksum and published with the
            tables as ``run_record.json``.

    Returns:
        (table name, rows written) for each table.
    """
    staged_tables = tables.copy()
    staged_report = deepcopy(report)
    out.parent.mkdir(parents=True, exist_ok=True)
    scratch = Path(tempfile.mkdtemp(prefix=".hww-", dir=out.parent))
    stage, backup = scratch / "tables", scratch / "previous"
    try:
        if out.exists():
            shutil.copytree(out, stage)
        else:
            stage.mkdir()
        counts = []
        for name, df in staged_tables.items():
            counts.append((name, write_csv(df, stage, name)))
            staged_tables[name] = df.reindex(columns=COLUMNS[name])
        add_output_counts(staged_report, counts, engagement)
        staged_report.write(stage / "data_quality_report.md", run_at)
        if "survey_other_text" not in staged_tables:
            (stage / "survey_other_text.csv").unlink(missing_ok=True)
        if record is not None:
            outputs: dict[str, dict[str, Any]] = {
                name: {"rows": n, "sha256": sha256_file(stage / f"{name}.csv")}
                for name, n in counts
            }
            outputs["data_quality_report"] = {
                "sha256": sha256_file(stage / "data_quality_report.md")
            }
            write_record(
                stage / RECORD_NAME,
                {
                    **record,
                    "outcome": "published",
                    "finished_at": utc_iso(),
                    "outputs": outputs,
                },
            )
        if out.exists():
            out.rename(backup)
        try:
            stage.rename(out)
        except OSError:
            if backup.exists():
                backup.rename(out)
            raise
        if backup.exists():
            shutil.rmtree(backup)
    finally:
        # Retain the old generation if restoration itself fails.
        if not backup.exists():
            shutil.rmtree(scratch)
    tables.update(staged_tables)
    report.sections = staged_report.sections
    report.unmapped = staged_report.unmapped
    return counts


# ---------------------------------------------------------------- main


def run(
    config_path: str | Path = DEFAULT_CONFIG,
    output_folder: str | Path | None = None,
    now: datetime | None = None,
    run_id: str | None = None,
) -> dict[str, pd.DataFrame]:
    """Run the whole pipeline.

    Args:
        config_path: The config file. Folder paths inside it are relative
            to it.
        output_folder: Write here instead of the output_folder in the
            config (the tests use a temporary folder).
        now: Treat this as the time of the run (the tests fix it).
            Defaults to the real time. It is shown in the report, and
            dates after its day are reported as in the future.
        run_id: The run's id, shared with its log (``main`` passes it);
            a new one by default.

    Returns:
        The output tables as a dict of DataFrames, keyed by table name.

    Raises:
        InputError: A config setting, mapping file, input file, sheet or
            column is missing or unreadable, or the outputs can't be
            written.
    """
    run_at = datetime.now() if now is None else now
    run_id = run_id or new_run_id()
    today = run_at.date()
    config_path = Path(config_path).resolve()
    cfg = load_config(config_path)
    base = config_path.parent
    inputs = (base / cfg["paths"]["input_folder"]).resolve()
    maps = (base / cfg["paths"]["mappings_folder"]).resolve()
    out = (
        Path(output_folder)
        if output_folder
        else (base / cfg["paths"]["output_folder"]).resolve()
    )
    for label, folder in (("input_folder", inputs), ("mappings_folder", maps)):
        if not folder.is_dir():
            raise InputError(
                f"{config_path.name}: {label} {folder} does not exist."
            )
    logger.info("Config: %s", config_path)
    output = cfg.get("output", {})
    # An output folder given by the caller (the tests, the golden
    # refresh) always means CSV files there.
    target = "folder" if output_folder else output.get("target", "folder")
    if target == "folder":
        try:
            out.mkdir(parents=True, exist_ok=True)
        except OSError as error:
            raise InputError(
                f"Cannot create the output folder {out} "
                f"({error.strerror or type(error).__name__}). Check the "
                "output_folder setting and that the folder can be written."
            ) from error
    bq_target = connect_bigquery(output) if target == "bigquery" else None
    report = Report()
    lookups, overrides = load_mappings(maps)
    record = start_record(
        run_id, run_at, config_path, sorted(maps.glob("*.csv"))
    )

    # 1. Read and clean each source.
    signups = load_signups(
        cfg["signups"], inputs, lookups, overrides, report, today
    )
    survey, answers, other_text, columns_map = load_survey(
        cfg["survey"],
        inputs,
        maps,
        lookups,
        overrides,
        report,
        cfg.get("privacy", {}).get("personal_columns", []),
    )
    was, sessions, topics = load_was(
        cfg["was"], inputs, lookups, overrides, report
    )
    peer = load_peer(cfg["peer"], inputs, lookups, overrides, report)

    # 2. One row per employer across all sources.
    org_rows = combine_org_rows(
        {"signups": signups, "survey": survey, "was": was, "peer": peer}
    )
    dim_employer = build_dim_employer(
        org_rows, cfg["employers"]["source_priority"], report
    )
    report_duplicates(signups, dim_employer, report)

    # 3. Attach employer_id to every source row.
    attach_employer_ids((signups, survey, was, peer, sessions), dim_employer)

    # 4. Survey tables.
    responses, fact_answers, catalogue = build_survey_tables(
        survey, answers, columns_map, dim_employer, maps, report
    )

    # 5. Engagement events.
    engagement = build_engagement(
        signups,
        survey,
        was,
        sessions,
        peer,
        cfg["was"]["completed_rule"],
        report,
        today,
    )
    reconcile(signups, survey, was, report)

    # 6. Calendar.
    dim_date = build_dim_date(
        *calendar_dates(engagement, fact_answers, sessions, signups)
    )

    # 7. Write everything.
    tables = {
        "dim_employer": dim_employer,
        "fact_signup": signups,
        "fact_engagement": engagement,
        "dim_survey_response": responses,
        "fact_survey_answer": fact_answers,
        "question_catalogue": catalogue,
        "fact_was_session": sessions,
        "fact_was_session_topic": topics,
        # No source yet for events and badges: written empty, with their
        # headers.
        "fact_event": pd.DataFrame(columns=COLUMNS["fact_event"]),
        "badges": pd.DataFrame(columns=COLUMNS["badges"]),
        "dim_date": dim_date,
    }
    if cfg["survey"]["keep_other_free_text"]:
        tables["survey_other_text"] = other_text
    checks = cfg.get("checks", {})
    compare_with_previous(
        report,
        previous_record(bq_target, out),
        float(checks.get("max_row_drop", 0.1)),
    )
    record["inputs"] = report.inputs
    record["attention"] = report.to_act_on()
    if bq_target is not None:
        counts = publish_bigquery(
            tables, engagement, report, run_at, record, bq_target
        )
        logger.info(
            "Published %d tables, the report and the run record to "
            "BigQuery dataset %s",
            len(tables),
            bq_target.dataset,
        )
    else:
        try:
            counts = write_outputs(
                tables, out, engagement, report, run_at, record
            )
        except OSError as error:
            raise InputError(
                f"Could not write the outputs to {out} "
                f"({error.strerror or type(error).__name__}). The previous "
                "outputs were kept. Close any output file that is open in "
                "another program, such as Excel, check the folder can be "
                "written, and run again."
            ) from error
        logger.info(
            "Wrote %d tables and data_quality_report.md to %s",
            len(tables),
            out,
        )
    for name, n in counts:
        logger.info("  %-24s %6d rows", name, n)
    attention = report.to_act_on()
    if attention:
        logger.warning(
            "Report needs attention: %s. See data_quality_report.md.",
            ", ".join(f"{n} {ATTENTION[k][0]}" for k, n in attention.items()),
        )
    else:
        logger.info("Report: nothing needs attention.")
    return tables


def main(argv: Sequence[str] | None = None) -> int:
    """Command-line entry point: ``python wrangle.py [config.toml]``.

    Args:
        argv: Arguments after the script name (defaults to
            ``sys.argv[1:]``).

    Returns:
        Exit code: 0 when the run finished; 1 when a problem a person
        can fix stopped it (an input, setting or mapping, or an output
        file open in another program), and the message says what to
        do; 2 for an unexpected error, which is a bug to report.
    """
    args = list(sys.argv[1:] if argv is None else argv)
    config_path = Path(args[0]) if args else DEFAULT_CONFIG
    try:
        mode = log_format()
    except ValueError as error:
        print(f"Stopped: {error}", file=sys.stderr)
        return 1
    run_id = new_run_id()
    folder = log_folder(config_path) if mode == "text" else None
    log_file = folder / f"{run_id}.log" if folder else None
    try:
        handlers = configure_logging(run_id, mode, log_file)
    except OSError:
        handlers = configure_logging(run_id, mode)
        logger.warning("Cannot write the log file %s; console only.", log_file)
        folder = log_file = None
    try:
        logger.info("Run %s", run_id)
        run(config_path, run_id=run_id)
    except InputError as error:
        logger.error("Stopped: %s", error)
        return 1
    except Exception as error:
        details = f" The details are in {log_file}." if log_file else ""
        logger.exception(
            "Unexpected error (%s at %s). This is a bug: please report it.%s",
            type(error).__name__,
            error_location(error),
            details,
        )
        return 2
    finally:
        close_logging(handlers)
        if folder is not None:
            prune_logs(folder)
    return 0


def log_folder(config_path: str | Path) -> Path | None:
    """Find the folder for run logs: [paths] log_folder, or logs/.

    Args:
        config_path: The config file. The folder is relative to it.

    Returns:
        The log folder, or None if the config can't be read (the run
        then stops and says why, on the console).
    """
    path = Path(config_path).resolve()
    try:
        cfg = load_config(path)
    except InputError:
        return None
    folder = str(cfg["paths"].get("log_folder", "logs"))
    return (path.parent / folder).resolve()


def error_location(error: BaseException) -> str:
    """Name the line of this package where an error was raised.

    Args:
        error: The exception, with its traceback.

    Returns:
        ``file.py:line`` of the innermost frame in hww_dashboard, or of
        the innermost frame at all, or ``"unknown"``.
    """
    package = Path(__file__).resolve().parent
    frames = traceback.extract_tb(error.__traceback__)
    ours = [
        f for f in frames if Path(f.filename).resolve().is_relative_to(package)
    ]
    if not frames:
        return "unknown"
    frame = (ours or frames)[-1]
    return f"{Path(frame.filename).name}:{frame.lineno}"


if __name__ == "__main__":
    sys.exit(main())
