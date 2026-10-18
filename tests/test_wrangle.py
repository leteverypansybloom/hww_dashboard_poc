"""Checks on the wrangling pipeline, grouped by the function under test."""

import re
import tomllib
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from hww_dashboard import wrangle

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
IP = re.compile(r"\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b")
LINK = re.compile(r"https?://")


def read_outputs(folder: Path) -> dict[str, pd.DataFrame]:
    """Read every output CSV in a folder as text.

    Args:
        folder: A pipeline output folder.

    Returns:
        Table name -> table, with blank cells as ''.
    """
    return {
        p.stem: pd.read_csv(p, dtype=str, keep_default_na=False)
        for p in folder.glob("*.csv")
    }


@pytest.fixture(scope="module")
def outputs(pipeline_output: Path) -> dict[str, pd.DataFrame]:
    """The tables written by one run on the synthetic inputs.

    Args:
        pipeline_output: Shared output from the synthetic pipeline run.

    Returns:
        Table name -> table, as text.
    """
    return read_outputs(pipeline_output)


@pytest.fixture(scope="module")
def settings(fixture_config: Path) -> wrangle.Config:
    """The frozen config that the run used.

    Args:
        fixture_config: Frozen configuration pointing only to synthetic inputs.

    Returns:
        tests/data/config.toml as nested dicts.
    """
    with open(fixture_config, "rb") as f:
        return tomllib.load(f)


@pytest.fixture(scope="module")
def input_cells(data_folder: Path) -> list[str]:
    """Every non-blank cell of every synthetic input workbook, as text.

    Args:
        data_folder: Folder containing synthetic workbooks and frozen settings.

    Returns:
        The cells' values, headings included.
    """
    found: list[str] = []
    for path in sorted((data_folder / "inputs").glob("*.xlsx")):
        book = load_workbook(path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                found.extend(str(v) for v in row if v is not None)
        book.close()
    return found


@pytest.fixture
def unpivoted() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Three responses to a made-up wide export, unpivoted.

    Q04 is tick-box: R1 ticks Minor illness, Mental health and Other and
    types free text in the Other box, R2 ticks Minor illness and R3
    ticks nothing. Q14 is single-choice: the answers are Yes, Weithiau
    (Welsh for Sometimes) and No.

    Returns:
        (answers, other_text) from ``unpivot_survey``, without free text.
    """
    wide = pd.DataFrame(
        {
            "response_id": ["R1", "R2", "R3"],
            "Minor illness": ["Minor illness", "Minor illness", ""],
            "Mental health": ["Mental health", "", ""],
            "Other": ["Other", "", ""],
            "Other.1": ["my free text", "", ""],
            "Flexible?": ["Yes", "Weithiau", "No"],
        }
    )
    columns_map = pd.DataFrame(
        [
            ["Minor illness", "tick", "Q04", "Minor illness"],
            ["Mental health", "tick", "Q04", "Mental health"],
            ["Other", "tick", "Q04", "Other"],
            ["Other.1", "other_text", "Q04", ""],
            ["Flexible?", "single", "Q14", ""],
        ],
        columns=["column", "role", "question_id", "option_or_note"],
    )
    answer_lookup = {
        ("Q14", "yes"): "Yes",
        ("Q14", "sometimes"): "Sometimes",
        ("Q14", "no"): "No",
    }
    translations = {("Q14", "weithiau"): "Sometimes"}
    answers, other_text, _ = wrangle.unpivot_survey(
        wide, columns_map, answer_lookup, translations, False
    )
    return answers, other_text


class TestRun:
    """``run``: privacy and key rules hold across the whole output."""

    def test_no_personal_columns_in_output(
        self, outputs: dict[str, pd.DataFrame], settings: wrangle.Config
    ) -> None:
        """No column named in [privacy] personal_columns reaches a table."""
        personal = {
            c.casefold() for c in settings["privacy"]["personal_columns"]
        }
        for name, df in outputs.items():
            found = [c for c in df.columns if c.casefold() in personal]
            assert not found, f"{name}.csv has personal column(s) {found}"

    def test_no_personal_values_in_output(
        self, outputs: dict[str, pd.DataFrame]
    ) -> None:
        """No email addresses, IP addresses or web links in any cell."""
        patterns = (
            (EMAIL, "an email"),
            (IP, "an IP address"),
            (LINK, "a link"),
        )
        for name, df in outputs.items():
            for column in df.columns:
                for value in df[column]:
                    for pattern, what in patterns:
                        assert not pattern.search(
                            value
                        ), f"{name}.{column} contains {what}: {value}"

    def test_every_employer_id_exists_in_dim_employer(
        self, outputs: dict[str, pd.DataFrame]
    ) -> None:
        """employer_id is unique in dim_employer; tables use known ids."""
        ids = set(outputs["dim_employer"]["employer_id"])
        rows = len(outputs["dim_employer"])
        assert len(ids) == rows, "employer_id is not unique in dim_employer"
        for name, df in outputs.items():
            if name == "dim_employer" or "employer_id" not in df.columns:
                continue
            unknown = set(df["employer_id"]) - ids - {""}
            assert not unknown, (
                f"{name}.csv has employer_id(s) missing from dim_employer: "
                f"{sorted(unknown)[:5]}"
            )

    def test_dim_employer_has_one_row_per_name(
        self, outputs: dict[str, pd.DataFrame]
    ) -> None:
        """Names that differ only by case are merged into one employer."""
        names = outputs["dim_employer"]["employer_name"].str.casefold()
        assert not names.duplicated().any()

    def test_every_signup_row_is_kept(
        self,
        outputs: dict[str, pd.DataFrame],
        settings: wrangle.Config,
        fixture_config: Path,
    ) -> None:
        """Every sign-up row stays a user, though employers are merged."""
        s = settings["signups"]
        inputs = fixture_config.parent / settings["paths"]["input_folder"]
        raw = wrangle.read_sheet(
            inputs / s["file"], s["sheet"], s["header_row"]
        )
        expected = (raw[s["org"]].map(wrangle.clean_text) != "").sum()
        events = outputs["fact_engagement"]["event_type"]
        assert len(outputs["fact_signup"]) == expected
        assert (events == "account_signup").sum() == expected

    def test_every_survey_answer_is_in_the_catalogue(
        self, outputs: dict[str, pd.DataFrame]
    ) -> None:
        """Every answer row links to a question_catalogue row."""
        answers = outputs["fact_survey_answer"]
        assert (answers["answer_id"] != "").all()

    def test_at_most_one_tick_per_respondent_and_option(
        self, outputs: dict[str, pd.DataFrame]
    ) -> None:
        """So counting rows in fact_survey_answer gives the tick count."""
        answers = outputs["fact_survey_answer"]
        key = ["response_id", "question_id", "answer"]
        assert not answers.duplicated(key).any()

    def test_report_uses_frozen_clock(self, pipeline_output: Path) -> None:
        """Shared pipeline reports carry only the injected timestamp."""
        report = (pipeline_output / "data_quality_report.md").read_text()
        assert "Run at 2026-10-10 12:00." in report


class TestSyntheticInputs:
    """The committed inputs hold only reserved, synthetic personal data."""

    def test_personal_values_are_present(self, input_cells: list[str]) -> None:
        """Emails, IPs and links exist, so the privacy tests can fail."""
        assert any(EMAIL.search(cell) for cell in input_cells)
        assert any(IP.search(cell) for cell in input_cells)
        assert any(LINK.search(cell) for cell in input_cells)

    def test_personal_values_are_reserved_examples(
        self, input_cells: list[str]
    ) -> None:
        """Every email, IP and link uses a reserved example name or range.

        example.org is reserved by RFC 2606 and 192.0.2.0/24 by RFC 5737,
        so these values can never reach a real person or site.
        """
        for cell in input_cells:
            for email in EMAIL.findall(cell):
                assert email.endswith("@example.org"), email
            for ip in IP.findall(cell):
                assert ip.startswith("192.0.2."), ip
            for site in re.findall(r"https?://([^/\s]+)", cell):
                assert site == "example.org", site


class TestUnpivotSurvey:
    """``unpivot_survey``: the wide export becomes one row per answer."""

    def test_free_text_stays_out(
        self, unpivoted: tuple[pd.DataFrame, pd.DataFrame]
    ) -> None:
        """Text typed into an Other box is neither an answer nor kept."""
        answers, other_text = unpivoted
        assert "my free text" not in set(answers["answer"])
        assert other_text.empty

    def test_ticked_boxes_give_the_mapped_option(
        self, unpivoted: tuple[pd.DataFrame, pd.DataFrame]
    ) -> None:
        """Any value in a tick-box column is the option from the mapping."""
        answers, _ = unpivoted
        r1 = answers[answers["response_id"] == "R1"]
        q04 = set(r1.loc[r1["question_id"] == "Q04", "answer"])
        assert q04 == {"Minor illness", "Mental health", "Other"}

    def test_welsh_answer_is_translated(
        self, unpivoted: tuple[pd.DataFrame, pd.DataFrame]
    ) -> None:
        """Weithiau is recorded as the English answer Sometimes."""
        answers, _ = unpivoted
        q14 = answers[answers["question_id"] == "Q14"]
        assert list(q14["answer"]) == ["Yes", "Sometimes", "No"]

    @pytest.mark.parametrize(
        ("role", "qid", "option"),
        [("tci k", "Q01", "A"), ("tick", "", "A"), ("tick", "Q01", "")],
    )
    def test_rejects_invalid_mapping(
        self, role: str, qid: str, option: str
    ) -> None:
        """Reject mapping typos rather than silently dropping answers.

        Args:
            role: Mapping role, possibly misspelt.
            qid: Question identifier, possibly blank.
            option: Tick option, possibly blank.
        """
        specs = pd.DataFrame(
            [["A", role, qid, option]],
            columns=wrangle.MAPPING_COLUMNS["survey_columns"],
        )
        with pytest.raises(wrangle.InputError, match="survey_columns.csv"):
            wrangle.unpivot_survey(
                pd.DataFrame({"response_id": ["R1"], "A": ["yes"]}),
                specs,
                {},
                {},
                False,
            )

    def test_tick_uses_catalogue_answer_case(self) -> None:
        """Case differences in tick labels cannot break catalogue links."""
        specs = pd.DataFrame(
            [["A", "tick", "Q01", "yes"]],
            columns=wrangle.MAPPING_COLUMNS["survey_columns"],
        )
        answers, _, _ = wrangle.unpivot_survey(
            pd.DataFrame({"response_id": ["R1"], "A": ["tick"]}),
            specs,
            {("Q01", "yes"): "Yes"},
            {},
            False,
        )
        assert list(answers["answer"]) == ["Yes"]

    def test_explicit_free_text_opt_in(self) -> None:
        """Opting in returns Other text separately from coded answers."""
        specs = pd.DataFrame(
            [["Other", "other_text", "Q01", ""]],
            columns=wrangle.MAPPING_COLUMNS["survey_columns"],
        )
        answers, other, new = wrangle.unpivot_survey(
            pd.DataFrame(
                {"response_id": ["R1"], "Other": ["Business suggestion"]}
            ),
            specs,
            {},
            {},
            True,
        )
        assert answers.empty
        assert not new
        assert other.to_dict("records") == [
            {
                "response_id": "R1",
                "question_id": "Q01",
                "other_text": "Business suggestion",
            }
        ]


class TestSurveyPercentages:
    """``survey_percentages``: Figure 13 numbers for one question."""

    def test_percent_is_ticks_over_respondents(
        self, unpivoted: tuple[pd.DataFrame, pd.DataFrame]
    ) -> None:
        """Percent = ticks / people who answered the question.

        R3 ticked nothing for Q04, so only R1 and R2 count as
        respondents.
        """
        answers, _ = unpivoted
        q4 = wrangle.survey_percentages(answers, "Q04").set_index("answer")
        assert (q4["respondents"] == 2).all()
        assert q4.loc["Minor illness", "ticks"] == 2
        assert q4.loc["Minor illness", "percent"] == pytest.approx(100.0)
        assert q4.loc["Mental health", "percent"] == pytest.approx(50.0)
        assert q4.loc["Other", "percent"] == pytest.approx(50.0)
        assert (q4["percent"] == q4["ticks"] / q4["respondents"] * 100).all()

    def test_single_choice_counts_everyone_who_answered(
        self, unpivoted: tuple[pd.DataFrame, pd.DataFrame]
    ) -> None:
        """All three responses answered Q14; one chose Sometimes."""
        answers, _ = unpivoted
        q14 = wrangle.survey_percentages(answers, "Q14").set_index("answer")
        assert q14.loc["Sometimes", "ticks"] == 1
        assert q14["respondents"].iloc[0] == 3


class TestCountValues:
    """``count_values``: counts in an order that every pandas gives."""

    def test_most_common_first_then_order_of_appearance(self) -> None:
        """Tied counts keep the order in which the values first appear."""
        values = pd.Series(["b", "c", "a", "c", "a", "b", "d", "c"])
        counts = wrangle.count_values(values)
        assert list(counts.items()) == [("c", 3), ("b", 2), ("a", 2), ("d", 1)]

    def test_blanks_are_not_counted(self) -> None:
        """Missing values are left out, as value_counts leaves them out."""
        values = pd.Series(["x", None, "x", float("nan")], dtype=object)
        assert list(wrangle.count_values(values).items()) == [("x", 2)]


class TestParseFlag:
    """``parse_flag``: 'Y (Dec)' style text on Account signups."""

    @pytest.mark.parametrize(
        ("text", "signup_month", "flag", "dates"),
        [
            ("", date(2025, 11, 1), False, []),
            ("Y", date(2025, 11, 1), True, []),
            ("\nY (Dec)", date(2025, 11, 1), True, [date(2025, 12, 1)]),
            ("Y (Dec)", date(2026, 1, 1), True, [date(2025, 12, 1)]),
            ("Y (Jan)", date(2025, 12, 1), True, [date(2026, 1, 1)]),
            (
                "Y (Nov 25 & Aug 26)",
                date(2025, 11, 1),
                True,
                [date(2025, 11, 1), date(2026, 8, 1)],
            ),
            ("Y (March 2026)", date(2026, 2, 1), True, [date(2026, 3, 1)]),
            ("Y (Sept 26)", date(2026, 7, 1), True, [date(2026, 9, 1)]),
            ("maybe", date(2026, 7, 1), None, []),
        ],
    )
    def test_reads_flag_and_dates(
        self,
        text: str,
        signup_month: date,
        flag: bool | None,
        dates: list[date],
    ) -> None:
        """The text gives a flag and first-of-month dates.

        A month without a year gets the year closest to the sign-up
        month.

        Args:
            text: The cell text.
            signup_month: The row's sign-up month.
            flag: The expected flag (None when unreadable).
            dates: The expected dates.
        """
        got_flag, got_dates, _ = wrangle.parse_flag(text, signup_month)
        assert got_flag == flag
        assert got_dates == dates


class TestCleanText:
    """``clean_text``: any cell becomes tidy text."""

    def test_numbers_and_text_of_one_value_match(self) -> None:
        """1, '1' and 1.0 all become '1', so lookups match them."""
        assert (
            wrangle.clean_text(1)
            == wrangle.clean_text("1")
            == wrangle.clean_text(1.0)
            == "1"
        )

    def test_stray_newlines_and_excel_codes_go(self) -> None:
        """Newlines and Excel's _x000D_ become single spaces, trimmed."""
        assert wrangle.clean_text("\n2-5") == "2-5"
        assert wrangle.clean_text("Sleep_x000D_\nGambling") == "Sleep Gambling"

    def test_blanks_become_empty_text(self) -> None:
        """None and NaN both become ''."""
        assert (
            wrangle.clean_text(None) == wrangle.clean_text(float("nan")) == ""
        )

    def test_nullable_scalar_is_blank(self) -> None:
        """Pandas nullable scalars cannot become literal <NA> identifiers."""
        assert wrangle.clean_text(pd.NA) == ""


class TestToDate:
    """Read Excel and text dates without ambiguity."""

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            ("2026-03-04", date(2026, 3, 4)),
            ("04/03/2026", date(2026, 3, 4)),
            (46085, date(2026, 3, 4)),
            (46085.5, date(2026, 3, 4)),
            (pd.NaT, None),
            (pd.NA, None),
            (True, None),
            (float("inf"), None),
            ("2026-02-30", None),
            ("", None),
        ],
    )
    def test_cell_types(self, value: object, expected: date | None) -> None:
        """Read supported dates and reject missing or invalid values.

        Args:
            value: An Excel cell or text date.
            expected: The exact calendar date, or no date.
        """
        assert wrangle.to_date(value) == expected


class TestAttachEmployerIds:
    """Attach employer keys even when a source contains no records."""

    def test_empty_frame_has_key_column(self) -> None:
        """Downstream joins require the key on header-only sources."""
        frame = pd.DataFrame(columns=["org"])
        dim = pd.DataFrame(columns=wrangle.COLUMNS["dim_employer"])
        wrangle.attach_employer_ids([frame], dim)
        assert list(frame.columns) == ["org", "employer_id"]
        assert frame.empty

    def test_case_matching_and_blank_orgs(self) -> None:
        """Known names link without case and unnamed responses stay blank."""
        frame = pd.DataFrame({"org": ["ORG A", "", "Missing"]})
        dim = pd.DataFrame({"employer_name": ["Org A"], "employer_id": ["E1"]})
        wrangle.attach_employer_ids([frame], dim)
        assert list(frame["employer_id"]) == ["E1", "", ""]


class TestMakeUnique:
    """Repeated Excel headers remain unique beside literal suffixes."""

    def test_reserves_existing_suffixes(self) -> None:
        """Generated names cannot overwrite an existing suffixed column."""
        assert wrangle.make_unique(["A", "A", "A.1", "A"]) == [
            "A",
            "A.2",
            "A.1",
            "A.3",
        ]


class TestWriteOutputs:
    """A failed generation preserves the last complete output set."""

    @pytest.mark.parametrize("failure", ["csv", "report", "publish"])
    def test_failure_keeps_previous_generation(
        self,
        tmp_path: Path,
        run_time: datetime,
        monkeypatch: pytest.MonkeyPatch,
        failure: str,
    ) -> None:
        """Inject failures in rendering and publication.

        Args:
            tmp_path: Temporary output location.
            run_time: Frozen pipeline clock.
            monkeypatch: Scoped failure injection.
            failure: Stage at which the run fails.
        """
        out = tmp_path / "out"
        out.mkdir()
        for name in (
            "badges.csv",
            "fact_event.csv",
            "data_quality_report.md",
            "survey_other_text.csv",
            "notes.txt",
        ):
            (out / name).write_text("old " + name, encoding="utf-8")
        before = {p.name: p.read_bytes() for p in out.iterdir()}
        tables = {
            "badges": pd.DataFrame({"badge": ["new"]}),
            "fact_event": pd.DataFrame(columns=wrangle.COLUMNS["fact_event"]),
        }
        original = wrangle.write_csv
        rename = Path.rename

        def write(df: pd.DataFrame, folder: Path, name: str) -> int:
            """Fail while writing the second table.

            Args:
                df: Output frame.
                folder: Staging folder.
                name: Table to write.

            Returns:
                Number of rows if successful.
            """
            if name == "fact_event":
                raise OSError("disk full")
            return original(df, folder, name)

        def fail_report(
            self: wrangle.Report, path: Path, run_at: datetime
        ) -> None:
            """Simulate a report write failure.

            Args:
                self: Report being rendered.
                path: Staged report path.
                run_at: Frozen run timestamp.
            """
            raise OSError("disk full")

        def fail_publish(self: Path, target: str | Path) -> Path:
            """Fail only when installing the staged generation.

            Args:
                self: Directory being renamed.
                target: Destination directory.

            Returns:
                The renamed path when restoring the old generation.
            """
            if self.name == "tables":
                raise OSError("publish failed")
            return rename(self, target)

        if failure == "csv":
            monkeypatch.setattr(wrangle, "write_csv", write)
        elif failure == "report":
            monkeypatch.setattr(wrangle.Report, "write", fail_report)
        else:
            monkeypatch.setattr(Path, "rename", fail_publish)
        report = wrangle.Report()
        with pytest.raises(OSError):
            wrangle.write_outputs(
                tables, out, pd.DataFrame({"event_type": []}), report, run_time
            )
        assert {p.name: p.read_bytes() for p in out.iterdir()} == before
        assert list(tables["badges"].columns) == ["badge"]
        assert not report.sections["Outputs"]

    def test_success_replaces_generation(
        self, tmp_path: Path, run_time: datetime
    ) -> None:
        """Successful publication removes stale free text and keeps notes."""
        out = tmp_path / "out"
        out.mkdir()
        (out / "notes.txt").write_text("keep", encoding="utf-8")
        (out / "survey_other_text.csv").write_text("stale", encoding="utf-8")
        tables = {"badges": pd.DataFrame({"badge": ["Award"]})}
        report = wrangle.Report()
        counts = wrangle.write_outputs(
            tables, out, pd.DataFrame({"event_type": []}), report, run_time
        )
        assert counts == [("badges", 1)]
        assert list(tables["badges"].columns) == wrangle.COLUMNS["badges"]
        assert (out / "notes.txt").read_text() == "keep"
        assert not (out / "survey_other_text.csv").exists()
        assert (
            "Run at 2026-10-10 12:00."
            in (out / "data_quality_report.md").read_text()
        )
        assert list(tmp_path.iterdir()) == [out]


class TestWriteCsv:
    """Output schema filtering also checks retained values for identifiers."""

    @pytest.mark.parametrize(
        "value", ["person@example.org", "192.0.2.4", "https://example.org/p"]
    )
    def test_blocks_personal_values(self, tmp_path: Path, value: str) -> None:
        """Stop personal text pasted into a retained output field.

        Args:
            tmp_path: Temporary output folder.
            value: Reserved synthetic personal identifier.
        """
        with pytest.raises(wrangle.InputError, match="badges.csv") as error:
            wrangle.write_csv(
                pd.DataFrame({"badge": [value]}), tmp_path, "badges"
            )
        assert value not in str(error.value)
        assert not (tmp_path / "badges.csv").exists()

    def test_schema_dates_and_private_columns(self, tmp_path: Path) -> None:
        """Extra personal columns are dropped and dates use ISO format."""
        frame = pd.DataFrame(
            {
                "employer_id": ["E1"],
                "badge": ["Award"],
                "awarded_date": [date(2026, 3, 4)],
                "Email": ["person@example.org"],
            }
        )
        assert wrangle.write_csv(frame, tmp_path, "badges") == 1
        written = pd.read_csv(tmp_path / "badges.csv", dtype=str)
        assert list(written.columns) == wrangle.COLUMNS["badges"]
        assert written.iloc[0].tolist() == ["E1", "Award", "2026-03-04"]


class TestReport:
    """Reports must not leak personal text through warnings."""

    def test_blocks_personal_values(
        self, tmp_path: Path, run_time: datetime
    ) -> None:
        """Reject identifiers in report notes without echoing them."""
        report = wrangle.Report()
        report.add("Warnings", "bad value person@example.org")
        path = tmp_path / "data_quality_report.md"
        with pytest.raises(wrangle.InputError, match="data_quality_report.md"):
            report.write(path, run_time)
        assert not path.exists()

    def test_repeated_write_is_identical(
        self, tmp_path: Path, run_time: datetime
    ) -> None:
        """Rendering a report twice must not duplicate unmapped notes."""
        report = wrangle.Report()
        report.unmapped[("sector", "signups", "New band")] = 2
        path = tmp_path / "report.md"
        report.write(path, run_time)
        first = path.read_bytes()
        report.write(path, run_time)
        assert path.read_bytes() == first
        assert report.sections["Unmapped values"] == []

    def test_zero_drops_add_no_note(self) -> None:
        """Report row drops only when at least one row is removed."""
        report = wrangle.Report()
        report.dropped("Survey", 0, "blank")
        assert report.sections["Rows dropped"] == []
        report.dropped("Survey", 2, "blank")
        assert report.sections["Rows dropped"] == [
            "- Survey: 2 row(s) dropped - blank"
        ]


class TestBuildQuestionCatalogue:
    """Catalogue keys must not drop or multiply survey answers."""

    @pytest.mark.parametrize("problem", ["unknown", "questions", "options"])
    def test_rejects_invalid_keys(self, tmp_path: Path, problem: str) -> None:
        """Reject orphan and duplicate catalogue keys before joining.

        Args:
            tmp_path: Temporary mapping folder.
            problem: Mapping defect to introduce.
        """
        question = "Q01,Question,Short,single\n"
        option = "Q01,Yes,Yes,1\n"
        (tmp_path / "survey_questions.csv").write_text(
            "question_id,question,question_short,question_type\n"
            + question
            + (question if problem == "questions" else ""),
            encoding="utf-8",
        )
        (tmp_path / "survey_answers.csv").write_text(
            "question_id,answer,answer_short,answer_order\n"
            + option
            + (option if problem == "options" else ""),
            encoding="utf-8",
        )
        specs = pd.DataFrame(
            [
                [
                    "Choice",
                    "tick",
                    "Q99" if problem == "unknown" else "Q01",
                    "Yes",
                ]
            ],
            columns=wrangle.MAPPING_COLUMNS["survey_columns"],
        )
        answers = pd.DataFrame(columns=["question_id", "answer"])
        with pytest.raises(wrangle.InputError, match="survey_.*csv"):
            wrangle.build_question_catalogue(tmp_path, specs, answers)


class TestMarkdownTable:
    """Render report tables without breaking Markdown columns."""

    def test_empty_and_pipes(self) -> None:
        """Empty tables use None and cell pipes become slashes."""
        assert wrangle.markdown_table(pd.DataFrame()) == "None."
        assert wrangle.markdown_table(pd.DataFrame({"A": ["x|y"]})) == (
            "| A |\n|---|\n| x/y |"
        )


class TestMapValue:
    """Map bands and count unknown values exactly once per cell."""

    def test_known_blank_and_unmapped(self) -> None:
        """Blank bands become Unknown while unknown topics remain text."""
        report = wrangle.Report()
        assert (
            wrangle.map_value(" A ", {"a": "Agreed"}, "d", "s", report)
            == "Agreed"
        )
        assert wrangle.map_value(None, {}, "d", "s", report) == "Unknown"
        assert wrangle.map_value("New", {}, "d", "s", report, True) == "New"
        assert wrangle.map_value("New", {}, "d", "s", report) == "Unknown"
        assert report.unmapped == {("d", "s", "New"): 2}


class TestMapBands:
    """Add mapped attributes in place without copying personal columns."""

    def test_maps_one_band(self) -> None:
        """Two raw cells give one known band and one Unknown."""
        frame = pd.DataFrame({"Raw": [" A ", None]})
        wrangle.map_bands(
            frame,
            {"sector": "Raw"},
            {"sector": {"a": "Public"}},
            "s",
            wrangle.Report(),
        )
        assert list(frame["sector"]) == ["Public", "Unknown"]


class TestOrgName:
    """Organisation overrides match cleaned, case-folded names."""

    def test_override_and_blank(self) -> None:
        """Spacing does not stop an override and blanks stay blank."""
        assert wrangle.org_name("  ORG  A ", {"org a": "Org B"}) == "Org B"
        assert wrangle.org_name(None, {}) == ""


class TestEmployerIdFor:
    """Employer keys are stable under case differences."""

    def test_known_hash_and_case(self) -> None:
        """The empty-string MD5 provides a fixed compatibility example."""
        assert wrangle.employer_id_for("") == "Ed41d8cd9"
        assert wrangle.employer_id_for("Org A") == wrangle.employer_id_for(
            "ORG A"
        )
        assert wrangle.employer_id_for("Org A") != wrangle.employer_id_for(
            "Org B"
        )


class TestIso:
    """Format optional dates without textual missing values."""

    def test_date_and_none(self) -> None:
        """Only a real date produces a formatted string."""
        assert wrangle.iso(date(2026, 3, 4)) == "2026-03-04"
        assert wrangle.iso(None) == ""


class TestFinYear:
    """The financial year changes on 1 April."""

    def test_year_boundary(self) -> None:
        """31 March and 1 April belong to adjacent financial years."""
        assert wrangle.fin_year(date(2026, 3, 31)) == "2025-26"
        assert wrangle.fin_year(date(2026, 4, 1)) == "2026-27"


class TestNearestYear:
    """Resolve yearless flag months relative to the sign-up month."""

    def test_exact_tie_prefers_later_year(self) -> None:
        """July is six months from both January candidates."""
        assert wrangle.nearest_year(1, date(2026, 7, 1)) == 2027


class TestPickValue:
    """Resolve employer attributes using each documented tie breaker."""

    @pytest.mark.parametrize(
        ("values", "sources", "positions", "expected"),
        [
            (["Unknown"], ["signups"], [0], "Unknown"),
            (["A", "A", "B"], ["peer", "peer", "signups"], [0, 1, 2], "A"),
            (["Other", "B"], ["signups", "peer"], [0, 1], "B"),
            (["A", "B"], ["peer", "signups"], [2, 1], "B"),
            (["A", "B"], ["signups", "signups"], [0, 1], "B"),
        ],
    )
    def test_ranking(
        self,
        values: list[str],
        sources: list[str],
        positions: list[int],
        expected: str,
    ) -> None:
        """Choose by frequency, specificity, source, then row position.

        Args:
            values: Candidate attributes including Unknown or Other.
            sources: Source of each candidate.
            positions: Combined input positions.
            expected: Winning attribute.
        """
        chosen, distinct = wrangle.pick_value(
            values, sources, positions, wrangle.SOURCES
        )
        assert chosen == expected
        assert distinct == sorted(set(values) - {"Unknown"})


class TestCombineOrgRows:
    """Stack source attributes in an explicit order."""

    def test_positions_follow_source_order(self) -> None:
        """Two employers receive positions zero and one in source order."""
        attrs = ["org", "sector", "industry", "size_band", "sites_band"]
        frame = pd.DataFrame(
            [["Org A", "Public", "Other", "0-9", "1"]], columns=attrs
        )
        combined = wrangle.combine_org_rows({"peer": frame, "was": frame})
        assert list(combined["source"]) == ["peer", "was"]
        assert list(combined["position"]) == [0, 1]
        assert list(combined["org"]) == ["Org A", "Org A"]


class TestBuildDimEmployer:
    """Merge case variants and preserve attribute conflict evidence."""

    def test_merges_names_and_reports_conflict(self) -> None:
        """A tied band chooses the later row within the same source."""
        frame = pd.DataFrame(
            {
                "org": ["Org A", "ORG A", ""],
                "sector": ["Public", "Private", "Unknown"],
                "industry": ["Unknown"] * 3,
                "size_band": ["Unknown"] * 3,
                "sites_band": ["Unknown"] * 3,
                "source": ["signups"] * 3,
                "position": [0, 1, 2],
            }
        )
        report = wrangle.Report()
        dim = wrangle.build_dim_employer(frame, wrangle.SOURCES, report)
        assert len(dim) == 1
        assert dim.iloc[0]["employer_name"] == "Org A"
        assert dim.iloc[0]["sector"] == "Private"
        assert dim.iloc[0]["attribute_conflicts"] == "sector: Private | Public"
        assert dim.iloc[0]["in_signups"]
        assert not dim.iloc[0]["in_was"]


class TestBuildDimDate:
    """Calendar bounds cover complete April-to-March financial years."""

    def test_bounds_leap_day_and_empty(self) -> None:
        """A leap day and following April span two full financial years."""
        frame = wrangle.build_dim_date(
            [date(2024, 2, 29), None], [date(2024, 4, 1)]
        )
        assert frame["date"].iloc[0] == date(2023, 4, 1)
        assert frame["date"].iloc[-1] == date(2025, 3, 31)
        assert len(frame) == 731
        april = frame[frame["date"] == date(2024, 4, 1)].iloc[0]
        assert april["fin_quarter"] == "Q1"
        assert april["fin_year"] == "2024-25"
        empty = wrangle.build_dim_date([None, pd.NaT])
        assert empty.empty
        assert list(empty.columns) == wrangle.COLUMNS["dim_date"]


class TestCalendarDates:
    """Calendar inputs include dates recorded only in summaries or flags."""

    def test_includes_summary_and_flags(self) -> None:
        """All four date sources contribute their literal dates."""
        day = date(2026, 3, 4)
        result = wrangle.calendar_dates(
            pd.DataFrame({"date": [day]}),
            pd.DataFrame({"date": [None]}),
            pd.DataFrame({"session_date": [None], "summary_sent_date": [day]}),
            pd.DataFrame(
                {"survey_dates": ["2026-01-01;2026-02-01"], "was_dates": [""]}
            ),
        )
        assert result == [
            [day],
            [None],
            [None, day],
            [date(2026, 1, 1), date(2026, 2, 1)],
        ]


class TestReconcile:
    """Reconciliation counts employers rather than rows or blank keys."""

    def test_counts_overlap_and_missing_flags(self) -> None:
        """Repeated signups count once and None is not a positive flag."""
        signups = pd.DataFrame(
            {
                "employer_id": ["E1", "E1", "E2"],
                "survey_flag": [True, True, None],
                "was_flag": [False, False, None],
            }
        )
        report = wrangle.Report()
        wrangle.reconcile(
            signups,
            pd.DataFrame({"employer_id": ["E1", "E2", ""]}),
            pd.DataFrame({"employer_id": []}),
            report,
        )
        assert (
            "1 employer(s) flagged 'Y'" in report.sections["Reconciliation"][0]
        )
        assert (
            "1 in both, 0 flagged with no record, 1 with a record but no flag"
            in report.sections["Reconciliation"][0]
        )


class TestUnpivotWasSessions:
    """Optional session details work independently of held dates."""

    def test_optional_columns_and_topic_deduplication(self) -> None:
        """A topic-only session is kept and repeated topics count once."""
        frame = pd.DataFrame(
            {
                "was_request_id": ["WAS-1"],
                "org": ["Org A"],
                "D": [None],
                "T": [" A;#A;#B"],
            }
        )
        cfg: wrangle.Config = {
            "sessions": [
                {"number": 1, "date": "D", "topics": "T"},
                {"number": 2, "date": "D"},
            ],
            "topic_separator": ";#",
        }
        sessions, topics = wrangle.unpivot_was_sessions(
            frame,
            cfg,
            {"was_topic": {"a": "Alpha", "b": "Beta"}},
            wrangle.Report(),
        )
        assert list(sessions["session_id"]) == ["WAS-1-S1"]
        assert sessions.iloc[0]["session_date"] is None
        assert sessions.iloc[0]["duration"] == ""
        assert list(topics["topic"]) == ["Alpha", "Beta"]


class TestCheckWasSessions:
    """Warn about undated details and impossible session chronology."""

    def test_undated_and_early_sessions(self) -> None:
        """Both conditions produce distinct actionable warnings."""
        was = pd.DataFrame(
            {"was_request_id": ["WAS-1"], "enquiry_date": [date(2026, 3, 4)]}
        )
        sessions = pd.DataFrame(
            {
                "was_request_id": ["WAS-1", "WAS-1"],
                "session_id": ["WAS-1-S1", "WAS-1-S2"],
                "session_number": [1, 2],
                "session_date": [None, date(2026, 3, 3)],
            }
        )
        report = wrangle.Report()
        wrangle.check_was_sessions(was, sessions, report)
        assert (
            "1 session row(s) have details but no date"
            in report.sections["Warnings"][0]
        )
        assert (
            "1 session(s) are dated before their enquiry date"
            in report.sections["Warnings"][1]
        )


class TestReportDuplicates:
    """Duplicate summaries use case-folded employer membership."""

    def test_signup_counts_and_overlap(self) -> None:
        """Two signups for one employer and a WAS-only employer are counted."""
        dim = pd.DataFrame(
            {
                "employer_name": ["Org A", "Org B"],
                "in_signups": [True, False],
                "in_survey": [True, False],
                "in_was": [False, True],
                "in_peer": [False, False],
            }
        )
        report = wrangle.Report()
        wrangle.report_duplicates(
            pd.DataFrame({"org": ["Org A", "ORG A"]}), dim, report
        )
        notes = "\n".join(report.sections["Duplicates"])
        assert "2 sign-up rows for 1 organisations" in notes
        assert "Org A: 2 rows" in notes
        assert "1 found in more than one source" in notes
        assert "['Org B']" in notes


class TestBuildSurveyTables:
    """Response attributes fall back to employers without losing answers."""

    def test_fills_unknown_attributes(self, project_folder: Path) -> None:
        """One known employer fills only Unknown response attributes."""
        survey = pd.DataFrame(
            {
                "response_id": ["R1"],
                "employer_id": ["E1"],
                "date": [date(2026, 3, 4)],
                "sector": ["Unknown"],
                "industry": ["Known"],
                "size_band": ["0-9"],
                "sites_band": ["1"],
            }
        )
        dim = pd.DataFrame(
            {
                "employer_id": ["E1"],
                "sector": ["Public"],
                "industry": ["Other"],
                "size_band": ["50-249"],
                "sites_band": ["6+"],
            }
        )
        answers = pd.DataFrame(
            {
                "response_id": ["R1"],
                "question_id": ["Q01"],
                "answer": ["Known answer"],
            }
        )
        specs = pd.DataFrame(columns=wrangle.MAPPING_COLUMNS["survey_columns"])
        responses, facts, catalogue = wrangle.build_survey_tables(
            survey,
            answers,
            specs,
            dim,
            project_folder / "mappings",
            wrangle.Report(),
        )
        assert responses.iloc[0]["sector"] == "Public"
        assert responses.iloc[0]["industry"] == "Known"
        assert len(facts) == 1
        assert facts.iloc[0]["answer_id"] in set(catalogue["answer_id"])


class TestBuildEngagement:
    """Each completion rule chooses its documented source date."""

    @pytest.mark.parametrize(
        ("rule", "expected"),
        [
            ("summary_1_sent", date(2026, 3, 5)),
            ("session_1_held", date(2026, 3, 4)),
            ("last_session_held", date(2026, 3, 6)),
        ],
    )
    def test_completion_rules(
        self, today: date, rule: str, expected: date
    ) -> None:
        """Build one held and one completed event per relevant date.

        Args:
            today: Frozen warning cutoff.
            rule: WAS completion configuration.
            expected: Expected completed event date.
        """
        empty = pd.DataFrame()
        sessions = pd.DataFrame(
            {
                "was_request_id": ["WAS-1"] * 2,
                "session_id": ["WAS-1-S1", "WAS-1-S2"],
                "employer_id": ["E1"] * 2,
                "session_number": [1, 2],
                "session_date": [date(2026, 3, 4), date(2026, 3, 6)],
                "summary_sent_date": [date(2026, 3, 5), None],
            }
        )
        fact = wrangle.build_engagement(
            empty, empty, empty, sessions, empty, rule, wrangle.Report(), today
        )
        assert list(
            fact.loc[fact["event_type"] == "was_completed", "date"]
        ) == [expected]
        assert len(fact) == 3

    def test_invalid_rule_with_no_sessions(self, today: date) -> None:
        """An invalid rule must fail even when there are no session rows."""
        empty = pd.DataFrame()
        sessions = pd.DataFrame(columns=wrangle.COLUMNS["fact_was_session"])
        with pytest.raises(wrangle.InputError, match="completed_rule"):
            wrangle.build_engagement(
                empty,
                empty,
                empty,
                sessions,
                empty,
                "typo",
                wrangle.Report(),
                today,
            )


class TestCompareWithPrevious:
    """``compare_with_previous``: far fewer rows than last time."""

    @staticmethod
    def report_with(rows: int) -> wrangle.Report:
        """Make a report whose WAS database had some rows read.

        Args:
            rows: Rows read this run.

        Returns:
            The report.
        """
        report = wrangle.Report()
        report.inputs.append({"source": "WAS database", "rows": rows})
        return report

    @pytest.mark.parametrize(
        ("rows", "previous", "warned"),
        [
            (90, 100, False),
            (89, 100, True),
            (150, 100, False),
            (0, 100, True),
            (50, 0, False),
        ],
    )
    def test_warns_past_the_threshold(
        self, rows: int, previous: int, warned: bool
    ) -> None:
        """A drop of more than 10% warns; growth never does.

        Args:
            rows: Rows read this run.
            previous: Rows read by the previous run.
            warned: Whether a warning is expected.
        """
        report = self.report_with(rows)
        record = {
            "run_id": "r1",
            "inputs": [{"source": "WAS database", "rows": previous}],
        }
        wrangle.compare_with_previous(report, record, 0.1)
        assert bool(report.sections["Warnings"]) is warned
        assert ("shrunk_sources" in report.to_act_on()) is warned

    def test_no_previous_run(self) -> None:
        """The first run has nothing to compare with."""
        report = self.report_with(1)
        wrangle.compare_with_previous(report, None, 0.1)
        assert report.sections["Warnings"] == []


class TestReadRecord:
    """``read_record``: the previous run's record, if there is one."""

    def test_missing_unreadable_and_valid(self, tmp_path: Path) -> None:
        """No file or a broken one gives None; a record is returned.

        Args:
            tmp_path: Holds the record files.
        """
        path = tmp_path / "run_record.json"
        assert wrangle.read_record(path) is None
        path.write_text("{broken", "utf-8")
        assert wrangle.read_record(path) is None
        path.write_text('{"run_id": "r1"}', "utf-8")
        assert wrangle.read_record(path) == {"run_id": "r1"}
