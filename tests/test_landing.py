"""Checks on ingestion into py-common's raw tables, and reading them back.

``python ingest.py`` lands the synthetic workbooks in local raw tables
(landing.py); ``wrangle.py`` with ``[landed] read = true`` reads them
back (landed.py). The end-to-end check runs both and compares every
table with the golden snapshot made from the workbooks directly.
"""

import json
import os
import shutil
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook

pytest.importorskip("py_common")

from py_common.contract import SheetRows  # noqa: E402
from py_common.errors import ServiceError  # noqa: E402
from py_common.model import Outcome, SourceItem  # noqa: E402
from py_common.raw import to_ndjson  # noqa: E402

from hww_dashboard import landed, landing, wrangle  # noqa: E402

GOLDEN = Path(__file__).resolve().parent / "golden"
WAS_FILE = "dummy_was_database.xlsx"


def landing_config(
    tmp_path: Path, data_folder: Path, read: bool, **landed_settings: str
) -> Path:
    """Write a config that lands and reads in tmp_path.

    Args:
        tmp_path: Folder for the config, landing folder, store, tables.
        data_folder: tests/data, whose config.toml is copied.
        read: The [landed] read setting.
        **landed_settings: More [landed] settings.

    Returns:
        The config file.
    """
    project = data_folder.parent.parent
    text = (data_folder / "config.toml").read_text(encoding="utf-8")
    inputs = (data_folder / "inputs").as_posix()
    maps = (project / "mappings").as_posix()
    text = text.replace(
        'input_folder = "inputs"', f'input_folder = "{inputs}"'
    )
    text = text.replace(
        'mappings_folder = "../../mappings"', f'mappings_folder = "{maps}"'
    )
    lines = [f"read = {'true' if read else 'false'}", 'target = "local"']
    lines += [f'{key} = "{value}"' for key, value in landed_settings.items()]
    text += "\n[landed]\n" + "\n".join(lines) + "\n"
    path = tmp_path / "config.toml"
    path.write_text(text, encoding="utf-8")
    return path


def newest_rows(folder: Path, table: str) -> list[dict[str, Any]]:
    """The rows of a local raw table's newest committed load.

    Args:
        folder: The tables folder.
        table: Table name, e.g. "hww_was".

    Returns:
        The rows, each with the raw table's columns.
    """
    record = landing.local_audit(folder, table)[-1]
    return landing.local_rows(folder, table, record)


def payload_keys(folder: Path, table: str) -> set[str]:
    """Every payload key in a local raw table's newest load.

    Args:
        folder: The tables folder.
        table: Table name.

    Returns:
        The keys.
    """
    return {
        key for row in newest_rows(folder, table) for key in row["payload"]
    }


def ndjson(sheets: dict[str, list[str]], version: str = "v1") -> bytes:
    """Raw NDJSON rows, one text column per sheet, as py-common writes it.

    Args:
        sheets: Sheet name -> the "name" value of each row.
        version: The upload's source version.

    Returns:
        The NDJSON.
    """
    data, _ = to_ndjson(
        {
            sheet: SheetRows(
                ["name"],
                [(n, (value,)) for n, value in enumerate(values, start=2)],
            )
            for sheet, values in sheets.items()
        },
        SourceItem("f.xlsx", "f.xlsx", version, 1),
        "abc",
    )
    return data


def settings_of(config: Path) -> wrangle.Config:
    """The config, loaded and checked as the pipeline loads it.

    Args:
        config: config.toml.

    Returns:
        Its settings.
    """
    return wrangle.load_config(config)


@pytest.fixture(scope="module")
def landed_run(
    tmp_path_factory: pytest.TempPathFactory, data_folder: Path
) -> Path:
    """Stage and land the synthetic workbooks once.

    Args:
        tmp_path_factory: For a folder shared by this module.
        data_folder: tests/data.

    Returns:
        The folder holding config.toml, landing/, store/ and landed/.
    """
    folder = tmp_path_factory.mktemp("landed")
    config = landing_config(folder, data_folder, read=True)
    results = landing.ingest(config, stage_from=data_folder / "inputs")
    text, ok = landing.summary(results)
    assert ok, text
    return folder


class TestContracts:
    """The HWW contracts accept the workbooks and keep personal data out."""

    def test_every_source_loads(self, landed_run: Path) -> None:
        """Each source's raw table has the workbook's data rows."""
        rows = {
            source: len(newest_rows(landed_run / "landed", table))
            for source, table in landing.TABLES.items()
        }
        assert rows == {"signups": 200, "survey": 103, "was": 35, "peer": 11}

    def test_one_upload_loads_signups_and_survey(
        self, landed_run: Path
    ) -> None:
        """Both tables come from the one monitoring workbook upload."""
        tables = landed.LocalTables(landed_run / "landed")
        signups = tables.newest_load("hww_signups")
        survey = tables.newest_load("hww_survey")
        assert signups.file == survey.file
        assert signups.file == "dummy_HWW monitoring and reporting.xlsx"
        assert signups.sha256 == survey.sha256

    def test_rows_say_where_they_came_from(self, landed_run: Path) -> None:
        """Each row has the raw table's columns and its Excel row."""
        first, *_ = newest_rows(landed_run / "landed", "hww_signups")
        assert tuple(first) == landing.RAW_COLUMNS
        assert first["sheet_name"] == "Account signups"
        assert (
            first["source_file"] == "dummy_HWW monitoring and reporting.xlsx"
        )
        assert first["row_number"] == 3

    def test_survey_info_loads_only_what_the_map_reads(
        self, landed_run: Path
    ) -> None:
        """Email address, Month and free text stay out; sites are in."""
        keys = payload_keys(landed_run / "landed", "hww_survey")
        assert "Email address" not in keys
        assert "Month" not in keys
        assert not [key for key in keys if "more info" in key]
        assert "Number of sites" in keys

    def test_new_survey_column_quarantines_the_whole_workbook(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """A new Survey info heading stops both sheets of the upload.

        Signups and survey come from one upload, so neither loads
        without the other.
        """
        config = landing_config(tmp_path, data_folder, read=False)
        cfg = settings_of(config)
        landing.stage(cfg, data_folder / "inputs", tmp_path / "landing")
        book = tmp_path / "landing" / "signups" / cfg["survey"]["file"]
        workbook = load_workbook(book)
        sheet = workbook["Survey info"]
        column = sheet.max_column + 1
        sheet.cell(2, column, "Phone number")
        sheet.cell(3, column, "Synthetic number")
        workbook.save(book)
        workbook.close()

        results = landing.ingest(config)

        (monitoring,) = results["signups"]
        assert monitoring.outcome == Outcome.QUARANTINED
        assert "Phone number" in "; ".join(monitoring.errors)
        assert not landing.local_audit(tmp_path / "landed", "hww_signups")
        assert not landing.local_audit(tmp_path / "landed", "hww_survey")

    @pytest.mark.parametrize("source", wrangle.SOURCES)
    def test_excluded_headings_are_not_loaded(
        self, source: str, landed_run: Path
    ) -> None:
        """No heading a source's contract excludes is a payload key.

        Args:
            source: The source checked.
            landed_run: The landed synthetic workbooks.
        """
        cfg = settings_of(landed_run / "config.toml")
        contract = landing.source_contracts(
            cfg, Path(cfg["paths"]["mappings_folder"])
        )[source]
        keys = payload_keys(landed_run / "landed", contract.table)
        assert contract.worksheet.excluded
        assert not keys & set(contract.worksheet.excluded)

    def test_personal_columns_are_not_loaded(self, landed_run: Path) -> None:
        """No column named in [privacy] reaches a raw table."""
        cfg = settings_of(landed_run / "config.toml")
        personal = set(cfg["privacy"]["personal_columns"])
        for table in landing.TABLES.values():
            keys = {
                landing.pandas_name(key)
                for key in payload_keys(landed_run / "landed", table)
            }
            assert not keys & personal, table

    def test_values_keep_their_excel_types(self, landed_run: Path) -> None:
        """Dates are ISO text and the WAS id a whole number."""
        rows = newest_rows(landed_run / "landed", "hww_was")
        dates = [
            row["payload"]["Date of Enquiry"]
            for row in rows
            if row["payload"]["Date of Enquiry"] is not None
        ]
        ids = [
            row["payload"]["Participant ID"]
            for row in rows
            if row["payload"]["Participant ID"] is not None
        ]
        assert dates and all(datetime.fromisoformat(d) for d in dates)
        assert ids and all(type(value) is int for value in ids)

    def test_every_config_column_is_in_its_contract(
        self, landed_run: Path
    ) -> None:
        """Each heading config.toml reads is a loaded column."""
        cfg = settings_of(landed_run / "config.toml")
        maps = Path(cfg["paths"]["mappings_folder"])
        contracts = landing.source_contracts(cfg, maps)
        for source in ("signups", "was", "peer"):
            known = set(contracts[source].names.values())
            needed = {
                value
                for key, value in cfg[source].items()
                if isinstance(value, str)
                and key not in ("file", "file_pattern", "sheet")
                and key not in ("completed_rule", "topic_separator")
            }
            if source == "was":
                needed |= {
                    value
                    for session in cfg["was"]["sessions"]
                    for key, value in session.items()
                    if key != "number"
                }
            assert needed <= known, (source, needed - known)

    def test_new_column_quarantines_the_file(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """A column the WAS contract doesn't know stops the load."""
        config = landing_config(tmp_path, data_folder, read=False)
        landing.stage(
            settings_of(config), data_folder / "inputs", tmp_path / "landing"
        )
        book = tmp_path / "landing" / "was" / WAS_FILE
        workbook = load_workbook(book)
        sheet = workbook.active
        assert sheet is not None
        sheet.cell(row=1, column=sheet.max_column + 1, value="Home address")
        sheet.cell(row=2, column=sheet.max_column, value="1 High Street")
        workbook.save(book)

        results = landing.ingest(config)

        (was,) = results["was"]
        assert was.outcome == Outcome.QUARANTINED
        assert "Home address" in "; ".join(was.errors)
        assert not (tmp_path / "landed" / "hww_was").exists()

    def test_wrong_type_quarantines_the_file(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """Text in a date column stops the load (schema checked)."""
        config = landing_config(tmp_path, data_folder, read=False)
        landing.stage(
            settings_of(config), data_folder / "inputs", tmp_path / "landing"
        )
        book = tmp_path / "landing" / "was" / WAS_FILE
        workbook = load_workbook(book)
        sheet = workbook.active
        assert sheet is not None
        sheet.cell(row=2, column=2, value="next Tuesday")
        workbook.save(book)

        (was,) = landing.ingest(config)["was"]

        assert was.outcome == Outcome.QUARANTINED

    def test_running_again_loads_nothing_twice(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """A second run with the same files skips them."""
        config = landing_config(tmp_path, data_folder, read=False)
        landing.ingest(config, stage_from=data_folder / "inputs")

        again = landing.ingest(config)

        assert all(
            r.outcome == Outcome.SKIPPED for rs in again.values() for r in rs
        )
        (record,) = landing.local_audit(tmp_path / "landed", "hww_peer")
        assert record["row_count"] == 11


class TestWorkbooks:
    """workbooks(): one contract and pipeline per landing folder."""

    def test_the_monitoring_workbook_holds_two_sheets(
        self, fixture_config: Path, project_folder: Path
    ) -> None:
        """Signups and survey share one contract, each with its table."""
        books = landing.workbooks(
            settings_of(fixture_config), project_folder / "mappings"
        )

        assert list(books) == ["signups", "was", "peer"]
        monitoring = books["signups"]
        assert monitoring.sources == ["signups", "survey"]
        assert monitoring.tables == {
            "Account signups": "hww_signups",
            "Survey info": "hww_survey",
        }
        assert monitoring.contract.file_name == (
            "dummy_HWW monitoring and reporting.xlsx"
        )
        assert set(monitoring.contract.untyped_columns) == {
            "Account signups",
            "Survey info",
        }


class TestRepeatedHeadings:
    """Headings convert between pandas' and py-common's numbering."""

    def test_round_trip(self) -> None:
        """Other.2 is py-common's Other_3, and back."""
        assert landing.py_common_name("Other.2") == "Other_3"
        assert landing.py_common_name("Other") == "Other"
        assert landing.pandas_name("col_B") == "B"
        assert landing.pandas_name("Organisation Industry ") == (
            "Organisation Industry"
        )


class TestLocalRawWarehouse:
    """The local stand-in follows RawJsonWarehouse's rules."""

    @staticmethod
    def warehouse(folder: Path) -> landing.LocalRawWarehouse:
        """A warehouse with two sheets, S and T.

        Args:
            folder: The tables folder.

        Returns:
            The warehouse.
        """
        return landing.LocalRawWarehouse(folder, {"S": "s_raw", "T": "t_raw"})

    def test_load_inserts_once_per_upload(self, tmp_path: Path) -> None:
        """A load inserts each sheet's rows; the same upload is skipped."""
        warehouse = self.warehouse(tmp_path)
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        data = ndjson({"S": ["a", "b"], "T": ["c"]})

        assert warehouse.load("f.xlsx_abc", "abc", data, 3, item)
        assert not warehouse.load("f.xlsx_abc", "abc", data, 3, item)

        assert len(newest_rows(tmp_path, "s_raw")) == 2
        assert len(newest_rows(tmp_path, "t_raw")) == 1
        assert [
            r["source_file"] for r in landing.local_audit(tmp_path, "s_raw")
        ] == ["f.xlsx"]

    def test_a_revert_loads_again(self, tmp_path: Path) -> None:
        """A, then B, then A again: each upload loads (latest wins)."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "s_raw"})
        uploads = [("v1", ["a"]), ("v2", ["a", "b"]), ("v3", ["a"])]

        loaded = [
            warehouse.load(
                "f.xlsx_abc",
                "abc",
                ndjson({"S": values}, version),
                len(values),
                SourceItem("f.xlsx", "f.xlsx", version, 1),
            )
            for version, values in uploads
        ]

        assert loaded == [True, True, True]
        audit = landing.local_audit(tmp_path, "s_raw")
        assert [r["source_version"] for r in audit] == ["v1", "v2", "v3"]

    def test_rows_have_the_bigquery_raw_columns(self, tmp_path: Path) -> None:
        """Local rows have exactly RawJsonWarehouse's raw table columns."""
        bigquery_raw = pytest.importorskip("py_common.adapters.bigquery_raw")
        warehouse = self.warehouse(tmp_path)
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        warehouse.load("k", "abc", ndjson({"S": ["a"]}), 1, item)

        (row,) = newest_rows(tmp_path, "s_raw")

        assert list(row) == [f.name for f in bigquery_raw.RAW_SCHEMA]
        assert list(landing.RAW_COLUMNS) == list(row)

    def test_a_table_added_later_loads_on_its_own(
        self, tmp_path: Path
    ) -> None:
        """Adding T to the mapping loads T, without loading S again."""
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        data = ndjson({"S": ["a"], "T": ["b"]})
        landing.LocalRawWarehouse(tmp_path, {"S": "s_raw"}).load(
            "k", "abc", ndjson({"S": ["a"]}), 1, item
        )

        assert self.warehouse(tmp_path).load("k", "abc", data, 2, item)

        assert len(landing.local_audit(tmp_path, "s_raw")) == 1
        assert len(landing.local_audit(tmp_path, "t_raw")) == 1

    def test_row_count_mismatch_is_refused(self, tmp_path: Path) -> None:
        """A wrong expected count fails the load, as in BigQuery."""
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        with pytest.raises(ValueError, match="Row count mismatch"):
            self.warehouse(tmp_path).load(
                "k", "abc", ndjson({"S": ["a"]}), 3, item
            )

    def test_a_sheet_with_no_table_is_refused(self, tmp_path: Path) -> None:
        """Rows from a sheet the mapping doesn't name are not loaded."""
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        with pytest.raises(ValueError, match="No raw table"):
            self.warehouse(tmp_path).load(
                "k", "abc", ndjson({"X": ["a"]}), 1, item
            )

    def test_identity_is_a_plain_name(self, tmp_path: Path) -> None:
        """Like RawJsonWarehouse's, with no colon or slash for keys."""
        identity = self.warehouse(tmp_path).identity
        assert identity.endswith(".s_raw+t_raw")
        assert ":" not in identity and "/" not in identity

    def test_a_failed_commit_is_retried_exactly_once(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Rows written without their audit record are not committed."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "s_raw"})
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        data = ndjson({"S": ["a", "b"]})
        real = os.replace

        def refuse_audit(source: Any, destination: Any) -> None:
            """Fail only the replace that would commit the audit record.

            Args:
                source: The temporary file.
                destination: The file it would replace.

            Raises:
                OSError: For the audit record.
            """
            if str(destination).endswith(".audit.json"):
                raise OSError("simulated audit failure")
            real(source, destination)

        with monkeypatch.context() as context:
            context.setattr("hww_dashboard.landing.os.replace", refuse_audit)
            with pytest.raises(ServiceError, match="audit failure"):
                warehouse.load("k", "abc", data, 2, item)

        assert landing.local_audit(tmp_path, "s_raw") == []
        assert warehouse.load("k", "abc", data, 2, item)
        assert not warehouse.load("k", "abc", data, 2, item)
        assert len(newest_rows(tmp_path, "s_raw")) == 2


class TestLongPath:
    """Kept copies get a path Windows' 260-character limit can't stop."""

    def test_windows_gets_the_extended_form(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """On Windows the path gets the extended prefix; elsewhere not."""
        monkeypatch.setattr("hww_dashboard.landing.os.name", "posix")
        assert landing.long_path(tmp_path) == tmp_path
        monkeypatch.setattr("hww_dashboard.landing.os.name", "nt")
        assert str(landing.long_path(tmp_path)).startswith("\\\\?\\")


class TestLandedInputs:
    """Reading the newest load back, shaped like the workbook."""

    def test_newest_load_wins(
        self, tmp_path: Path, data_folder: Path, run_time: datetime
    ) -> None:
        """After a second upload, the reader returns only its rows."""
        config = landing_config(tmp_path, data_folder, read=True)
        landing.ingest(config, stage_from=data_folder / "inputs")
        book = tmp_path / "landing" / "peer" / "dummy_peer_montoring.xlsx"
        workbook = load_workbook(book)
        sheet = workbook.active
        assert sheet is not None
        sheet.delete_rows(3, sheet.max_row)
        workbook.save(book)
        landing.ingest(config)

        cfg = settings_of(config)
        inputs = landed.landed_inputs(
            cfg, tmp_path, Path(cfg["paths"]["mappings_folder"])
        )
        sheet_input = inputs.read("peer", cfg["peer"], "Peer mentoring", [])

        assert len(sheet_input.frame) == 1
        assert sheet_input.facts["table"] == "hww_peer"
        assert sheet_input.name == "dummy_peer_montoring.xlsx"

    def test_frame_matches_the_workbook_reader(
        self, landed_run: Path, data_folder: Path
    ) -> None:
        """Headings, row numbers and values match read_sheet's."""
        cfg = settings_of(landed_run / "config.toml")
        inputs = landed.landed_inputs(
            cfg, landed_run, Path(cfg["paths"]["mappings_folder"])
        )
        from_table = inputs.read("was", cfg["was"], "WAS database", []).frame
        from_book = wrangle.read_sheet(
            data_folder / "inputs" / WAS_FILE, 0, 1, "WAS database"
        )
        columns = [c for c in from_table.columns]
        assert list(from_table["excel_row"]) == list(from_book["excel_row"])
        left = from_table[columns].map(wrangle.clean_text)
        right = from_book[columns].map(wrangle.clean_text)
        pd.testing.assert_frame_equal(left, right)

    def test_survey_frame_matches_the_workbook_reader(
        self, landed_run: Path, data_folder: Path
    ) -> None:
        """Survey info read back has read_sheet's headings and values."""
        cfg = settings_of(landed_run / "config.toml")
        inputs = landed.landed_inputs(
            cfg, landed_run, Path(cfg["paths"]["mappings_folder"])
        )
        from_table = inputs.read(
            "survey", cfg["survey"], "Survey info", []
        ).frame
        from_book = wrangle.read_sheet(
            data_folder / "inputs" / cfg["survey"]["file"], "Survey info", 2
        )
        columns = list(from_table.columns)
        assert "Other (more info).6" in from_table.attrs["excluded"]
        left = from_table[columns].map(wrangle.clean_text)
        right = from_book[columns].map(wrangle.clean_text)
        pd.testing.assert_frame_equal(left, right)

    def test_blank_ticks_load_as_not_ticked(
        self, landed_run: Path, data_folder: Path
    ) -> None:
        """A tick cell holding a space or a non-breaking space is blank."""
        cfg = settings_of(landed_run / "config.toml")
        inputs = landed.landed_inputs(
            cfg, landed_run, Path(cfg["paths"]["mappings_folder"])
        )
        frame = inputs.read("survey", cfg["survey"], "Survey info", []).frame
        book = wrangle.read_sheet(
            data_folder / "inputs" / cfg["survey"]["file"], "Survey info", 2
        )
        raw = list(book["Caring responsibilities"])
        assert " " in raw and "\xa0" in raw
        ticked = [
            v
            for v in frame["Caring responsibilities"]
            if wrangle.clean_text(v)
        ]
        assert ticked == [v for v in raw if isinstance(v, str) and v.strip()]

    def test_missing_table_says_to_run_ingest(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """Reading before ingest.py has run stops with a clear message."""
        config = landing_config(tmp_path, data_folder, read=True)
        with pytest.raises(wrangle.InputError, match="python ingest.py"):
            wrangle.run(config, output_folder=tmp_path / "out")


def fake_bigquery(
    audit: list[tuple[Any, ...]],
    rows: list[tuple[Any, ...]],
    queries: list[tuple[str, Any]] | None = None,
) -> SimpleNamespace:
    """A stand-in BigQuery client for BigQueryTables.

    Args:
        audit: Rows the audit query returns: (load key, checksum,
            timestamp, row count, source version).
        rows: Rows the table query returns: (row number, payload).
        queries: Where to record each (SQL, job config), if given.

    Returns:
        An object with the client's query() method.
    """

    def query(sql: str, job_config: Any = None) -> SimpleNamespace:
        """Answer the audit query or the raw table query.

        Args:
            sql: The query text.
            job_config: Its parameters.

        Returns:
            A job whose result() holds the rows.
        """
        if queries is not None:
            queries.append((sql, job_config))
        if "FROM `p.d.audit`" in sql:
            return SimpleNamespace(result=lambda: audit)
        return SimpleNamespace(result=lambda: rows)

    return SimpleNamespace(query=query)


class TestBigQueryTables:
    """The BigQuery reader's queries, against a fake client."""

    def test_reads_the_selected_load_and_its_audit_row(self) -> None:
        """Rows come from the audited upload; facts from the audit table."""
        loaded = datetime(2026, 10, 1, 9, 0, tzinfo=UTC)
        key = "hww/was/dummy_was_database.xlsx_abcdef12"
        queries: list[tuple[str, Any]] = []
        client = fake_bigquery(
            [(key, "abcdef1234", loaded, 1, "17")],
            [(2, '{"Participant ID": 7}')],
            queries,
        )

        load = landed.BigQueryTables("p", "d", client).newest_load("hww_was")

        assert load.file == "dummy_was_database.xlsx"
        assert load.sha256 == "abcdef1234"
        assert load.rows == [(2, {"Participant ID": 7})]
        audit_sql, _ = queries[0]
        assert "FROM `p.d.audit`" in audit_sql
        assert "destination = @destination" in audit_sql
        rows_sql, config = queries[1]
        assert "WHERE source_generation = @generation" in rows_sql
        assert "AND ingested_at = @loaded_at" in rows_sql
        assert [q.value for q in config.query_parameters] == ["17", loaded]


class TestCheckLanded:
    """[landed] settings are checked before anything runs."""

    @pytest.mark.parametrize(
        "section, message",
        [
            ({"read": "yes"}, "read must be true or false"),
            ({"target": "cloud"}, "target must be one of"),
            ({"target": "gcp"}, "project must be nonblank"),
            ({"tables_folder": ""}, "tables_folder must be nonblank"),
        ],
    )
    def test_bad_settings(self, section: dict[str, Any], message: str) -> None:
        """Each bad setting names itself."""
        with pytest.raises(wrangle.InputError, match=message):
            wrangle.check_landed(section, "config.toml")


class TestEndToEnd:
    """Workbooks -> py-common raw tables -> wrangle -> same tables."""

    def test_tables_match_the_golden_snapshot(
        self, landed_run: Path, tmp_path: Path, run_time: datetime
    ) -> None:
        """Every output CSV equals the one made from the workbooks."""
        out = tmp_path / "out"
        wrangle.run(
            landed_run / "config.toml", output_folder=out, now=run_time
        )

        for golden in sorted(GOLDEN.glob("*.csv")):
            expected = golden.read_bytes().replace(b"\r\n", b"\n")
            actual = (out / golden.name).read_bytes().replace(b"\r\n", b"\n")
            assert actual == expected, golden.name

    def test_run_record_names_the_landed_tables(
        self, landed_run: Path, tmp_path: Path, run_time: datetime
    ) -> None:
        """The run record says which table and load each input was."""
        out = tmp_path / "out"
        wrangle.run(
            landed_run / "config.toml", output_folder=out, now=run_time
        )

        record = json.loads((out / "run_record.json").read_text("utf-8"))
        tables = {
            entry["source"]: entry["table"] for entry in record["inputs"]
        }
        assert tables == {
            "Account signups": "hww_signups",
            "Survey info": "hww_survey",
            "WAS database": "hww_was",
            "Peer mentoring": "hww_peer",
        }


def test_stage_copies_each_named_workbook(
    tmp_path: Path, data_folder: Path
) -> None:
    """--stage puts each workbook in its landing folder, once."""
    config = landing_config(tmp_path, data_folder, read=False)
    copies = landing.stage(
        settings_of(config), data_folder / "inputs", tmp_path / "landing"
    )
    assert sorted(p.parent.name for p in copies) == ["peer", "signups", "was"]
    signups = tmp_path / "landing" / "signups"
    assert [p.name for p in signups.iterdir()] == [
        "dummy_HWW monitoring and reporting.xlsx"
    ]
    shutil.rmtree(tmp_path / "landing")


class TestLandingFolders:
    """landing_folders(): which folder each source reads."""

    def test_sources_sharing_a_workbook_share_a_folder(
        self, fixture_config: Path
    ) -> None:
        """Signups and survey both read the monitoring workbook's folder."""
        folders = landing.landing_folders(settings_of(fixture_config))
        assert folders == {
            "signups": "signups",
            "survey": "signups",
            "was": "was",
            "peer": "peer",
        }

    def test_separate_workbooks_get_their_own_folders(
        self, fixture_config: Path
    ) -> None:
        """A source with its own file, or a file_pattern, reads its own."""
        cfg = deepcopy(settings_of(fixture_config))
        cfg["survey"]["file"] = "survey.xlsx"
        cfg["peer"] = {**cfg["peer"], "file_pattern": "*peer*.xlsx"}
        folders = landing.landing_folders(cfg)
        assert folders["survey"] == "survey"
        assert folders["peer"] == "peer"


class TestLandingRegressions:
    """Failures found while reviewing the ingestion boundary."""

    def test_signups_rejects_unexpected_personal_column(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """An extra signup column never reaches its raw table."""
        config = landing_config(tmp_path, data_folder, read=True)
        cfg = settings_of(config)
        landing.stage(cfg, data_folder / "inputs", tmp_path / "landing")
        book = tmp_path / "landing" / "signups" / cfg["signups"]["file"]
        workbook = load_workbook(book)
        sheet = workbook["Account signups"]
        column = sheet.max_column + 1
        sheet.cell(2, column, "Home address")
        sheet.cell(3, column, "Synthetic address")
        workbook.save(book)
        workbook.close()

        (result,) = landing.ingest(config)["signups"]

        assert result.outcome == Outcome.QUARANTINED
        assert "Home address" in "; ".join(result.errors)
        assert not (tmp_path / "landed" / "hww_signups").exists()

    def test_signups_requires_the_mixed_type_column(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """Number of sites is required even without a cell type check."""
        config = landing_config(tmp_path, data_folder, read=True)
        cfg = settings_of(config)
        landing.stage(cfg, data_folder / "inputs", tmp_path / "landing")
        book = tmp_path / "landing" / "signups" / cfg["signups"]["file"]
        workbook = load_workbook(book)
        sheet = workbook["Account signups"]
        for cell in sheet[2]:
            if cell.value == "Number of sites":
                sheet.delete_cols(cell.column)
                break
        workbook.save(book)
        workbook.close()

        (result,) = landing.ingest(config)["signups"]

        assert result.outcome == Outcome.QUARANTINED
        assert "Number of sites" in "; ".join(result.errors)

    @pytest.mark.parametrize("source", wrangle.SOURCES)
    def test_allow_empty_is_per_source(
        self, source: str, fixture_config: Path, project_folder: Path
    ) -> None:
        """Allowing one empty source does not mutate other contracts."""
        cfg = deepcopy(settings_of(fixture_config))
        cfg[source]["allow_empty"] = True

        contracts = landing.source_contracts(cfg, project_folder / "mappings")

        assert contracts[source].worksheet.allow_empty
        defaults = landing.source_contracts(
            settings_of(fixture_config), project_folder / "mappings"
        )
        assert not defaults[source].worksheet.allow_empty

    def test_pattern_selects_newest_export_before_loading(
        self, tmp_path: Path, data_folder: Path
    ) -> None:
        """Filename dates win over discovery order and unrelated files.

        A source with file_pattern reads a folder named after itself.
        """
        config = landing_config(tmp_path, data_folder, read=True)
        workbook = settings_of(config)["survey"]["file"]
        pattern = "*healthy-working-wales-employer-survey-tool-*.xlsx"
        text = config.read_text(encoding="utf-8")
        survey_lines = f'file = "{workbook}"\nsheet = "Survey info"'
        assert text.count(survey_lines) == 1
        config.write_text(
            text.replace(
                survey_lines,
                f'file_pattern = "{pattern}"\nsheet = "Survey info"',
            ),
            encoding="utf-8",
        )
        folder = tmp_path / "landing" / "survey"
        folder.mkdir(parents=True)
        newest = "a-healthy-working-wales-employer-survey-tool-2026-10-01.xlsx"
        older = "z-healthy-working-wales-employer-survey-tool-2026-09-01.xlsx"
        for name in (newest, older, "unrelated-2026-12-01.xlsx"):
            shutil.copyfile(data_folder / "inputs" / workbook, folder / name)

        results = landing.ingest(config)["survey"]

        assert [result.item.name for result in results] == [newest]
        assert results[0].outcome == Outcome.LOADED
        assert (
            landed.LocalTables(tmp_path / "landed")
            .newest_load("hww_survey")
            .file
            == newest
        )

    def test_empty_load_replaces_previous_rows(self, tmp_path: Path) -> None:
        """An audited zero-row snapshot does not revive an older load."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "t"})
        warehouse.load(
            "old",
            "old",
            ndjson({"S": ["a", "b"]}),
            2,
            SourceItem("f.xlsx", "f.xlsx", "v1", 1),
        )
        warehouse.load(
            "empty",
            "empty",
            b"",
            0,
            SourceItem("f.xlsx", "f.xlsx", "v2", 1),
        )

        load = landed.LocalTables(tmp_path).newest_load("t")

        assert load.rows == []
        assert load.sha256 == "empty"

    def test_existing_survey_history_uses_filename_dates(
        self, tmp_path: Path
    ) -> None:
        """Older exports loaded last cannot win on load time alone."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "t"})
        data = ndjson({"S": ["a"]})
        newest = "a-survey-2026-10-01.xlsx"
        older = "z-survey-2026-09-01.xlsx"
        for name in (newest, older):
            warehouse.load(
                name, name, data, 1, SourceItem(name, name, "v1", 1)
            )

        load = landed.LocalTables(tmp_path).newest_load(
            "t", {"file_pattern": "*-survey-*.xlsx"}
        )

        assert load.file == newest

    def test_exact_filename_can_select_a_previous_committed_load(
        self, tmp_path: Path
    ) -> None:
        """Changing a pinned file does not silently read another file."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "t"})
        data = ndjson({"S": ["a"]})
        for name in ("wanted.xlsx", "other.xlsx"):
            warehouse.load(
                name, name, data, 1, SourceItem(name, name, "v1", 1)
            )

        load = landed.LocalTables(tmp_path).newest_load(
            "t", {"file": "wanted.xlsx"}
        )

        assert load.file == "wanted.xlsx"

    def test_bigquery_reads_empty_load_from_audit(self) -> None:
        """The newest committed audit row selects even an empty load."""
        loaded = datetime(2026, 10, 2, 9, 0)
        client = fake_bigquery(
            [("hww/was/f.xlsx_abcdef12", "abcdef1234", loaded, 0, "3")], []
        )

        load = landed.BigQueryTables("p", "d", client).newest_load("hww_was")

        assert load.rows == []
        assert load.loaded_at == loaded
        assert load.sha256 == "abcdef1234"

    def test_empty_peer_snapshot_reaches_dashboard_tables(
        self, tmp_path: Path, data_folder: Path, run_time: datetime
    ) -> None:
        """The configured empty-source path works end to end."""
        config = landing_config(tmp_path, data_folder, read=True)
        config.write_text(
            config.read_text(encoding="utf-8").replace(
                "[peer]", "[peer]\nallow_empty = true"
            ),
            encoding="utf-8",
        )
        cfg = settings_of(config)
        landing.ingest(config, stage_from=data_folder / "inputs")
        book = tmp_path / "landing" / "peer" / cfg["peer"]["file"]
        workbook = load_workbook(book)
        sheet = workbook["query (1)"]
        sheet.delete_rows(2, sheet.max_row)
        workbook.save(book)
        workbook.close()

        (result,) = landing.ingest(config)["peer"]

        assert result.outcome == Outcome.LOADED
        assert result.rows_processed == 0
        tables = wrangle.run(
            config, output_folder=tmp_path / "out", now=run_time
        )
        assert "peer_mentoring_interest" not in set(
            tables["fact_engagement"]["event_type"]
        )


class TestSourceSelection:
    """Local and cloud discovery use the same export ordering rules."""

    def test_same_date_uses_modification_time_and_ignores_lock_files(
        self,
    ) -> None:
        """Lock files and unrelated names cannot outrank valid exports."""
        names = [
            "a-survey-2026-10-01.xlsx",
            "z-survey-2026-10-01.xlsx",
            "~$survey-2026-12-01.xlsx",
            "unrelated-2026-12-01.xlsx",
        ]
        items = [SourceItem(name, name, "1", 1) for name in names]
        source = SimpleNamespace(
            items=lambda: reversed(items), download=lambda item: b"original"
        )
        times = {names[0]: 2.0, names[1]: 1.0}
        selected = landing.NewestSource(
            source, "*survey-*.xlsx", lambda item: times[item.name]
        )

        (item,) = selected.items()

        assert item.name == names[0]
        assert selected.download(item) == b"original"

    def test_no_matches_produces_no_source_items(self) -> None:
        """A nonmatching file never reaches contract validation."""
        source = SimpleNamespace(
            items=lambda: [SourceItem("other.xlsx", "other.xlsx", "1", 1)]
        )
        selected = landing.NewestSource(source, "survey-*.xlsx", lambda _: 0)

        assert list(selected.items()) == []

    def test_cloud_build_selects_the_newest_matching_generation(
        self,
        fixture_config: Path,
        project_folder: Path,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Cloud selection reads timestamps for discovered generations."""
        cfg = settings_of(fixture_config)
        cfg["survey"].pop("file")
        cfg["survey"]["file_pattern"] = "*survey-*.xlsx"
        names = ["a-survey-2026-10-01.xlsx", "z-survey-2026-09-01.xlsx"]
        items = [
            SourceItem(name, f"hww/survey/{name}", str(n), 1)
            for n, name in enumerate(names, 1)
        ]
        requested: list[tuple[str, int]] = []

        def get_blob(name: str, generation: int) -> SimpleNamespace:
            """Return metadata without contacting a cloud service."""
            requested.append((name, generation))
            return SimpleNamespace(updated=datetime(2026, 10, 2, tzinfo=UTC))

        source = SimpleNamespace(
            items=lambda: items,
            download=lambda item: b"generation bytes",
            bucket=SimpleNamespace(get_blob=get_blob),
        )
        monkeypatch.setattr(
            "py_common.adapters.gcs.GCSSource", lambda *args: source
        )
        monkeypatch.setattr(
            "py_common.adapters.gcs.GCSObjectStore",
            lambda *args: SimpleNamespace(),
        )
        monkeypatch.setattr(
            "py_common.adapters.bigquery_raw.RawJsonWarehouse",
            lambda *args: SimpleNamespace(),
        )
        settings = landing.LandingSettings(
            "gcp",
            tmp_path,
            tmp_path,
            tmp_path,
            project="p",
            dataset="d",
            landing_prefix="hww",
            landing_bucket="landing",
            store_bucket="store",
        )

        pipelines = landing.build_pipelines(
            cfg, project_folder / "mappings", settings
        )
        (selected,) = pipelines["survey"].source.items()

        assert selected == items[0]
        assert requested == [
            (item.identity, int(item.version)) for item in items
        ]
        assert (
            pipelines["survey"].source.download(selected)
            == b"generation bytes"
        )

    def test_missing_cloud_metadata_fails_closed(self) -> None:
        """A missing generation is not assigned a fabricated timestamp."""
        source = SimpleNamespace(
            bucket=SimpleNamespace(get_blob=lambda *args, **kwargs: None)
        )

        with pytest.raises(ServiceError, match="modification time"):
            landing.gcs_modified(
                source, SourceItem("f.xlsx", "hww/f.xlsx", "1", 1)
            )


class TestCommittedSnapshots:
    """Readers reject incomplete loads and preserve empty snapshots."""

    def test_first_local_load_can_be_empty(self, tmp_path: Path) -> None:
        """An audit record identifies a load with no rows."""
        warehouse = landing.LocalRawWarehouse(tmp_path, {"S": "t"})
        item = SourceItem("f.xlsx", "f.xlsx", "v1", 1)
        warehouse.load("empty", "empty", b"", 0, item)

        load = landed.LocalTables(tmp_path).newest_load("t")

        assert load.rows == []
        assert load.file == "f.xlsx"
        assert load.sha256 == "empty"
        assert not warehouse.load("empty", "empty", b"", 0, item)

    def test_missing_committed_cloud_load_does_not_query_data(self) -> None:
        """Rows without a successful audit cannot become current data."""

        def query(sql: str, job_config: Any = None) -> SimpleNamespace:
            """Return no committed loads and refuse any table query."""
            assert "FROM `p.d.audit`" in sql
            return SimpleNamespace(result=lambda: [])

        tables = landed.BigQueryTables("p", "d", SimpleNamespace(query=query))

        with pytest.raises(wrangle.InputError, match="no committed load"):
            tables.newest_load("hww_was")

    def test_cloud_history_selects_filename_date_before_load_time(
        self,
    ) -> None:
        """Historical uploads use the same filename policy as discovery."""
        newer = datetime(2026, 10, 2, tzinfo=UTC)
        older = datetime(2026, 10, 3, tzinfo=UTC)
        queries: list[tuple[str, Any]] = []
        client = fake_bigquery(
            [
                (
                    "z-survey-2026-09-01.xlsx_bbbbbbbb",
                    "bbbbbbbb",
                    older,
                    1,
                    "2",
                ),
                (
                    "a-survey-2026-10-01.xlsx_aaaaaaaa",
                    "aaaaaaaa",
                    newer,
                    1,
                    "1",
                ),
            ],
            [(2, {"value": 1})],
            queries,
        )

        load = landed.BigQueryTables("p", "d", client).newest_load(
            "t", {"file_pattern": "*survey-*.xlsx"}
        )

        assert load.file == "a-survey-2026-10-01.xlsx"
        assert load.sha256 == "aaaaaaaa"
        audit_sql, audit_config = queries[0]
        assert "LIMIT 1" not in audit_sql
        assert audit_config.query_parameters[0].value == "p.d.t"
        _, rows_config = queries[1]
        assert [q.value for q in rows_config.query_parameters] == ["1", newer]

    def test_cloud_row_count_must_match_committed_audit(self) -> None:
        """A partial or inconsistent snapshot cannot enter wrangling."""
        loaded = datetime(2026, 10, 2, tzinfo=UTC)
        client = fake_bigquery(
            [("f.xlsx_aaaaaaaa", "aaaaaaaa", loaded, 2, "1")],
            [(2, {"value": 1})],
        )

        tables = landed.BigQueryTables("p", "d", client)

        with pytest.raises(wrangle.InputError, match="row count differs"):
            tables.newest_load("t")
